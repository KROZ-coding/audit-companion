import asyncio
import json
import os
import unittest
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from threading import BoundedSemaphore, Lock
from time import sleep
import time
from unittest.mock import patch
from zipfile import ZipFile
import httpx
from openpyxl import load_workbook

# Tests must not pick up server/.env (real LLM key) or the real knowledge index.
os.environ["SKIP_ENV_FILE"] = "1"
os.environ["KNOWLEDGE_INDEX_PATH"] = "./data/knowledge/__disabled_in_tests__.pkl"


from fastapi.testclient import TestClient  # noqa: E402

from app.config import LLMProvider, Settings, _llm_providers  # noqa: E402
from app.main import create_app  # noqa: E402
from app.models import Course, GradingResult, QuizSession  # noqa: E402
from app.services.knowledge_base import LocalKnowledgeBase, build_index  # noqa: E402
from app.services.excel_export import run_snapshot_scheduler  # noqa: E402
from app.services.llm_client import LLMClient, _quota_for  # noqa: E402
from app.services.maxkb_client import MaxKBClient  # noqa: E402
from app.store import Store  # noqa: E402


class SmokeTests(unittest.TestCase):
    def setUp(self) -> None:
        self.app = create_app()
        self.client = TestClient(self.app)

    def test_health(self) -> None:
        response = self.client.get("/health")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["status"], "ok")
        self.assertEqual(self.client.get("/ready").status_code, 200)
        self.assertEqual(response.headers["x-content-type-options"], "nosniff")
        frontend = self.client.get("/")
        self.assertEqual(frontend.status_code, 200)
        self.assertIn("text/html", frontend.headers["content-type"])
        self.assertIn("经管数智审计智能学伴", frontend.text)
        self.assertEqual(self.client.get("/课程思政地图.html").status_code, 200)

    def test_login_quiz_submit_flow(self) -> None:
        login = self.client.post("/api/auth/login", json={"username": "stu001", "password": "stu123*"})
        self.assertEqual(login.status_code, 200)
        login_audit = next(item for item in self.app.state.store.audit_logs if item["action"] == "login_success")
        self.assertEqual(login_audit["ip"], "testclient")
        self.assertEqual(self.client.get("/api/auth/me").json()["role"], "student")
        self.assertEqual(self.client.get("/api/auth/me").json()["courses"][0]["id"], "audit-101")

        started = self.client.post(
            "/api/quiz/start",
            json={"counts": {"single": 1, "multi": 1, "short": 0}},
        )
        self.assertEqual(started.status_code, 200)
        questions = started.json()["questions"]
        self.assertNotIn("answer", questions[0])
        self.assertNotIn("reference_answer", questions[0])
        answers = {questions[0]["id"]: 1, questions[1]["id"]: [0, 1, 2, 3]}
        submitted = self.client.post(f"/api/quiz/{started.json()['id']}/submit", json={"answers": answers})
        self.assertEqual(submitted.status_code, 200)
        self.assertEqual(submitted.json()["result"]["status"], "graded")
        self.assertEqual(
            self.client.get(f"/api/grading/{started.json()['id']}/status").status_code,
            200,
        )
        duplicate = self.client.post(f"/api/quiz/{started.json()['id']}/submit", json={"answers": answers})
        self.assertEqual(duplicate.status_code, 409)
        self.assertEqual(self.app.state.store.audit_logs[-1]["action"], "quiz_submit_rejected")

    def test_quiz_chapter_filter(self) -> None:
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "stu001", "password": "stu123*"}
            ).status_code,
            200,
        )
        filtered = self.client.post(
            "/api/quiz/start",
            json={
                "course_id": "audit-101",
                "chapter": "审计证据",
                "counts": {"single": 1, "multi": 0, "short": 0},
            },
        )
        self.assertEqual(filtered.status_code, 200)
        self.assertEqual(filtered.json()["questions"][0]["chapter"], "审计证据")
        missing = self.client.post(
            "/api/quiz/start",
            json={
                "course_id": "audit-101",
                "chapter": "不存在的章节",
                "counts": {"single": 1, "multi": 0, "short": 0},
            },
        )
        self.assertEqual(missing.status_code, 409)

    def test_quiz_difficulty_filter(self) -> None:
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "stu001", "password": "stu123*"}
            ).status_code,
            200,
        )
        filtered = self.client.post(
            "/api/quiz/start",
            json={
                "course_id": "audit-101",
                "difficulty": "medium",
                "counts": {"single": 0, "multi": 1, "short": 0},
            },
        )
        self.assertEqual(filtered.status_code, 200)
        self.assertEqual(filtered.json()["questions"][0]["type"], "multi_choice")
        missing = self.client.post(
            "/api/quiz/start",
            json={
                "course_id": "audit-101",
                "difficulty": "hard",
                "counts": {"single": 1, "multi": 0, "short": 0},
            },
        )
        self.assertEqual(missing.status_code, 409)

    def test_progress_history_contains_wrong_question_summary_only(self) -> None:
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "stu001", "password": "stu123*"}
            ).status_code,
            200,
        )
        quiz = self.client.post(
            "/api/quiz/start",
            json={"course_id": "audit-101", "counts": {"single": 1, "multi": 0, "short": 0}},
        ).json()
        question_id = quiz["questions"][0]["id"]
        self.assertEqual(
            self.client.post(
                f"/api/quiz/{quiz['id']}/submit", json={"answers": {question_id: 0}}
            ).status_code,
            200,
        )
        history = self.client.get("/api/progress/history").json()
        row = next(item for item in history if item["session_id"] == quiz["id"])
        self.assertEqual(len(row["wrong_questions"]), 1)
        self.assertNotIn("reference_answer", row["wrong_questions"][0])

    def test_question_bank_filters(self) -> None:
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "teacher01", "password": "teach123*"}
            ).status_code,
            200,
        )
        filtered = self.client.get(
            "/api/bank",
            params={"course_id": "audit-101", "type": "multi_choice", "difficulty": "medium", "chapter": "审计程序"},
        )
        self.assertEqual(filtered.status_code, 200)
        self.assertEqual(len(filtered.json()), 1)
        self.assertEqual(filtered.json()[0]["type"], "multi_choice")
        self.assertEqual(
            self.client.get("/api/bank", params={"difficulty": "hard"}).json(),
            [],
        )

    def test_assigned_quiz_is_listed_and_loadable_by_student(self) -> None:
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "teacher01", "password": "teach123*"}
            ).status_code,
            200,
        )
        question = next(
            item for item in self.app.state.store.questions.values()
            if item.status == "published" and item.course_id == "audit-101"
        )
        student = self.app.state.store.students_for_course("audit-101")[0]
        due_at = "2027-01-15T12:00:00+00:00"
        assigned = self.client.post(
            "/api/quiz/assign",
            json={"course_id": "audit-101", "title": "第一章测验", "question_ids": [question.id], "student_ids": [student.id], "due_at": due_at},
        )
        self.assertEqual(assigned.status_code, 200)
        assignment_id = assigned.json()["assignment_id"]
        tracking = self.client.get("/api/quiz/assigned", params={"course_id": "audit-101"}).json()
        tracked = next(item for item in tracking if item["id"] == assignment_id)
        self.assertEqual(tracked["target_count"], 1)
        self.assertEqual(tracked["due_at"], due_at)
        self.assertEqual(tracked["roster"][0]["student_id"], student.id)
        self.client.post("/api/auth/logout")
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "stu001", "password": "stu123*"}
            ).status_code,
            200,
        )
        assignments = self.client.get("/api/quiz/assignments", params={"course_id": "audit-101"})
        self.assertEqual(assignments.status_code, 200)
        assignment = next(item for item in assignments.json() if item["title"] == "第一章测验")
        self.assertEqual(assignment["status"], "ongoing")
        self.assertEqual(assignment["due_at"], due_at)
        session = self.client.get(f"/api/quiz/{assignment['id']}")
        self.assertEqual(session.status_code, 200)
        self.assertEqual(session.json()["title"], "第一章测验")
        self.assertEqual(session.json()["questions"][0]["id"], question.id)

    def test_unassigned_submitted_quiz_can_be_resumed(self) -> None:
        self.client.post("/api/auth/login", json={"username": "stu001", "password": "stu123*"})
        started = self.client.post(
            "/api/quiz/start",
            json={"course_id": "audit-101", "counts": {"single": 1, "multi": 0, "judge": 0, "fill": 0, "short": 0, "case": 0}},
        ).json()
        session_id = started["id"]
        question_id = started["questions"][0]["id"]
        self.app.state.store.quiz_sessions[session_id].submitted_answers = {question_id: 2}
        self.app.state.store.save()

        resumable = self.client.get("/api/quiz/resumable", params={"course_id": "audit-101"})
        session = self.client.get(f"/api/quiz/{session_id}")

        self.assertEqual(resumable.status_code, 200)
        self.assertEqual([item["id"] for item in resumable.json()], [session_id])
        self.assertEqual(session.json()["saved_answers"], {question_id: 2})

    def test_quiz_submission_is_persisted_before_grading(self) -> None:
        from app.services.grading_service import grade_session

        with TemporaryDirectory() as data_dir:
            settings = Settings(app_env="development", data_dir=data_dir, persistence_enabled=True, cookie_secure=False)
            with patch("app.main.settings", settings):
                app = create_app()
                store = app.state.store
                with TestClient(app) as client:
                    client.post("/api/auth/login", json={"username": "stu001", "password": "stu123*"})
                    quiz = client.post(
                        "/api/quiz/start",
                        json={"course_id": "audit-101", "counts": {"single": 1, "multi": 0, "judge": 0, "fill": 0, "short": 0, "case": 0}},
                    ).json()
                    question_id = quiz["questions"][0]["id"]
                    answers = {question_id: 1}

                    def inspect_saved_submission(session, submitted, current_settings):
                        persisted = json.loads(store.persistence_path.read_text(encoding="utf-8"))
                        saved = persisted["quiz_sessions"][session.id]
                        self.assertEqual(saved["status"], "grading")
                        self.assertEqual(saved["submitted_answers"], answers)
                        return grade_session(session, submitted, current_settings)

                    with patch("app.routers.quiz.grade_session", side_effect=inspect_saved_submission):
                        response = client.post(f"/api/quiz/{quiz['id']}/submit", json={"answers": answers})

                    self.assertEqual(response.status_code, 200)
                    self.assertFalse(store.quiz_sessions[quiz["id"]].submitted_answers)

    def test_quiz_draft_round_trip_and_resume(self) -> None:
        self.client.post("/api/auth/login", json={"username": "stu001", "password": "stu123*"})
        started = self.client.post(
            "/api/quiz/start",
            json={"course_id": "audit-101", "counts": {"single": 1, "multi": 0, "judge": 0, "fill": 0, "short": 0, "case": 0}},
        ).json()
        session_id = started["id"]
        question_id = started["questions"][0]["id"]

        saved = self.client.post(f"/api/quiz/{session_id}/draft", json={"answers": {question_id: 2}})
        loaded = self.client.get(f"/api/quiz/{session_id}").json()
        resumable = self.client.get("/api/quiz/resumable", params={"course_id": "audit-101"})

        self.assertEqual(saved.status_code, 200)
        self.assertTrue(saved.json()["saved"])
        self.assertIsNotNone(saved.json()["draft_updated_at"])
        self.assertEqual(loaded["draft_answers"], {question_id: 2})
        self.assertIsNotNone(loaded["draft_updated_at"])
        self.assertEqual([item["id"] for item in resumable.json()], [session_id])

    def test_quiz_draft_survives_store_restart(self) -> None:
        with TemporaryDirectory() as data_dir:
            settings = Settings(app_env="development", data_dir=data_dir, persistence_enabled=True, cookie_secure=False)
            with patch("app.main.settings", settings):
                app = create_app()
                with TestClient(app) as client:
                    client.post("/api/auth/login", json={"username": "stu001", "password": "stu123*"})
                    quiz = client.post(
                        "/api/quiz/start",
                        json={"course_id": "audit-101", "counts": {"single": 1, "multi": 0, "judge": 0, "fill": 0, "short": 0, "case": 0}},
                    ).json()
                    question_id = quiz["questions"][0]["id"]
                    response = client.post(f"/api/quiz/{quiz['id']}/draft", json={"answers": {question_id: 1}})
                    self.assertEqual(response.status_code, 200)

            resumed = Store(settings).quiz_sessions[quiz["id"]]

            self.assertEqual(resumed.status, "ongoing")
            self.assertEqual(resumed.draft_answers, {question_id: 1})
            self.assertIsNotNone(resumed.draft_updated_at)

    def test_quiz_draft_unchanged_content_skips_disk_write(self) -> None:
        self.client.post("/api/auth/login", json={"username": "stu001", "password": "stu123*"})
        started = self.client.post(
            "/api/quiz/start",
            json={"course_id": "audit-101", "counts": {"single": 1, "multi": 0, "judge": 0, "fill": 0, "short": 0, "case": 0}},
        ).json()
        session_id = started["id"]
        question_id = started["questions"][0]["id"]
        payload = {"answers": {question_id: 1}}
        calls = 0
        original = self.app.state.store.save_async

        async def counting():
            nonlocal calls
            calls += 1
            return await original()

        first = self.client.post(f"/api/quiz/{session_id}/draft", json=payload)
        self.assertTrue(first.json()["saved"])

        with patch.object(self.app.state.store, "save_async", counting):
            second = self.client.post(f"/api/quiz/{session_id}/draft", json=payload)

        # 内容未变时端点自身不再写盘；窗口内的至多一次保存来自全局持久化中间件兜底。
        self.assertFalse(second.json()["saved"])
        self.assertEqual(second.json()["draft_updated_at"], first.json()["draft_updated_at"])
        self.assertLessEqual(calls, 1)

    def test_quiz_draft_is_cleared_after_successful_submit(self) -> None:
        self.client.post("/api/auth/login", json={"username": "stu001", "password": "stu123*"})
        started = self.client.post(
            "/api/quiz/start",
            json={"course_id": "audit-101", "counts": {"single": 1, "multi": 0, "judge": 0, "fill": 0, "short": 0, "case": 0}},
        ).json()
        session_id = started["id"]
        question_id = started["questions"][0]["id"]
        self.client.post(f"/api/quiz/{session_id}/draft", json={"answers": {question_id: 0}})
        answers = {question_id: 1}

        submitted = self.client.post(f"/api/quiz/{session_id}/submit", json={"answers": answers})
        session = self.client.get(f"/api/quiz/{session_id}").json()

        self.assertEqual(submitted.status_code, 200)
        self.assertEqual(session["status"], "graded")
        self.assertEqual(session["draft_answers"], {})
        self.assertIsNone(session["draft_updated_at"])

    def test_quiz_draft_rejections(self) -> None:
        self.client.post("/api/auth/login", json={"username": "stu001", "password": "stu123*"})
        started = self.client.post(
            "/api/quiz/start",
            json={"course_id": "audit-101", "counts": {"single": 1, "multi": 0, "judge": 0, "fill": 0, "short": 0, "case": 0}},
        ).json()
        session_id = started["id"]
        question_id = started["questions"][0]["id"]

        unknown = self.client.post(f"/api/quiz/{session_id}/draft", json={"answers": {"ghost": 1}})
        missing = self.client.post("/api/quiz/does-not-exist/draft", json={"answers": {}})

        self.assertEqual(unknown.status_code, 422)
        self.assertEqual(missing.status_code, 404)

        answers = {question_id: 1}
        submitted = self.client.post(f"/api/quiz/{session_id}/submit", json={"answers": answers})
        late = self.client.post(f"/api/quiz/{session_id}/draft", json={"answers": answers})

        self.assertEqual(submitted.status_code, 200)
        self.assertEqual(late.status_code, 409)
        self.assertEqual(self.app.state.store.audit_logs[-1]["action"], "quiz_submit")

    def test_judge_fill_and_case_questions(self) -> None:
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "teacher01", "password": "teach123*"}
            ).status_code,
            200,
        )
        questions = [
            {
                "type": "judge", "stem": "审计证据需要充分且适当。", "answer": True,
                "course_id": "audit-101", "knowledge_points": ["审计证据"],
            },
            {
                "type": "fill", "stem": "审计证据的数量要求是____。", "answer": "充分性",
                "reference_answer": "充分性", "course_id": "audit-101", "knowledge_points": ["审计证据"],
            },
            {
                "type": "case", "stem": "指出案例中的主要审计风险。", "answer": None,
                "reference_answer": "审计风险", "rubric": [{"point": "审计风险", "score": 30}],
                "course_id": "audit-101", "knowledge_points": ["审计风险"],
            },
        ]
        created_ids = []
        for payload in questions:
            response = self.client.post("/api/bank", json=payload)
            self.assertEqual(response.status_code, 201)
            created_ids.append(response.json()["id"])
        for question_id in created_ids:
            self.assertEqual(self.client.post(f"/api/bank/{question_id}/submit").status_code, 200)
            self.assertEqual(
                self.client.post(f"/api/bank/{question_id}/review", json={"decision": "pass"}).json()["status"],
                "published",
            )

        self.client.post("/api/auth/logout")
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "stu001", "password": "stu123*"}
            ).status_code,
            200,
        )
        quiz = self.client.post(
            "/api/quiz/start",
            json={"course_id": "audit-101", "counts": {"single": 0, "multi": 0, "judge": 1, "fill": 1, "short": 0, "case": 1}},
        )
        self.assertEqual(quiz.status_code, 200)
        answers = {
            question["id"]: True if question["type"] == "judge" else "充分性" if question["type"] == "fill" else "本案例存在审计风险。"
            for question in quiz.json()["questions"]
        }
        submitted = self.client.post(f"/api/quiz/{quiz.json()['id']}/submit", json={"answers": answers})
        self.assertEqual(submitted.status_code, 200)
        result = submitted.json()["result"]
        self.assertEqual(result["status"], "needs_review")
        self.assertEqual(result["total"], 50)
        case_id = next(question["id"] for question in quiz.json()["questions"] if question["type"] == "case")
        self.client.post("/api/auth/logout")
        self.client.post("/api/auth/login", json={"username": "teacher01", "password": "teach123*"})
        reviewed = self.client.post(
            f"/api/grading/{result['id']}/review",
            json={"question_id": case_id, "score": 30, "reason": "教师确认案例分析结论正确"},
        )
        self.assertEqual(reviewed.status_code, 200)
        self.assertEqual(reviewed.json()["status"], "graded")
        self.assertEqual(reviewed.json()["total"], 50)

    def test_login_lockout_and_logout(self) -> None:
        for attempt in range(4):
            response = self.client.post(
                "/api/auth/login", json={"username": "stu001", "password": "wrong"}
            )
            self.assertEqual(response.status_code, 401, attempt)

        locked = self.client.post(
            "/api/auth/login", json={"username": "stu001", "password": "wrong"}
        )
        self.assertEqual(locked.status_code, 429)
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "stu001", "password": "stu123*"}
            ).status_code,
            429,
        )
        self.assertEqual(self.app.state.store.audit_logs[-1]["ip"], "testclient")

    def test_password_verification_uses_bounded_slots(self) -> None:
        from app.utils.security import verify_password

        class FakeHasher:
            def __init__(self):
                self.active = 0
                self.peak = 0
                self.lock = Lock()

            def verify(self, _hash, _password):
                with self.lock:
                    self.active += 1
                    self.peak = max(self.peak, self.active)
                try:
                    sleep(0.01)
                    return True
                finally:
                    with self.lock:
                        self.active -= 1

        fake = FakeHasher()
        with patch("app.utils.security._hasher", fake), patch(
            "app.utils.security._hash_slots", BoundedSemaphore(3)
        ):
            with ThreadPoolExecutor(max_workers=16) as pool:
                results = list(pool.map(lambda _: verify_password("pw", "hash"), range(24)))

        self.assertTrue(all(results))
        self.assertEqual(fake.peak, 3)

    def test_role_and_course_boundaries(self) -> None:
        login = self.client.post(
            "/api/auth/login", json={"username": "stu001", "password": "stu123*"}
        )
        self.assertEqual(login.status_code, 200)
        self.assertEqual(self.client.get("/api/courses").json()[0]["id"], "audit-101")
        self.assertEqual(self.client.get("/api/courses/audit-101/students").status_code, 403)
        self.assertEqual(self.client.get("/api/bank").status_code, 403)

        quiz = self.client.post(
            "/api/quiz/start",
            json={"course_id": "audit-101", "counts": {"single": 1, "multi": 0, "short": 0}},
        )
        self.assertEqual(quiz.status_code, 200)
        question = quiz.json()["questions"][0]
        self.assertNotIn("answer", question)
        self.assertNotIn("reference_answer", question)
        session_id = quiz.json()["id"]
        self.assertEqual(
            self.client.post(
                f"/api/quiz/{session_id}/submit", json={"answers": {"outside-question": 1}}
            ).status_code,
            422,
        )

        self.assertEqual(self.client.post("/api/auth/logout").status_code, 204)
        self.assertEqual(self.client.get("/api/auth/me").status_code, 401)

    def test_teacher_and_admin_access(self) -> None:
        teacher = self.client.post(
            "/api/auth/login", json={"username": "teacher01", "password": "teach123*"}
        )
        self.assertEqual(teacher.status_code, 200)
        self.assertEqual(self.client.get("/api/courses/audit-101/students").status_code, 200)
        self.assertEqual(self.client.get("/api/bank").status_code, 200)
        self.assertEqual(self.client.get("/api/chat/monitor").status_code, 200)
        self.assertEqual(self.client.get("/api/admin/users").status_code, 403)
        self.client.post("/api/auth/logout")

        admin = self.client.post(
            "/api/auth/login", json={"username": "admin", "password": "admin123*"}
        )
        self.assertEqual(admin.status_code, 200)
        self.assertEqual(self.client.get("/api/admin/users").status_code, 200)
        logs = self.client.get("/api/admin/audit-logs").json()
        self.assertTrue(any(item["action"] == "login_success" for item in logs))

    def test_course_admin_and_enrollment_flow(self) -> None:
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "admin", "password": "admin123*"}
            ).status_code,
            200,
        )
        teacher_id = next(item.id for item in self.app.state.store.users.values() if item.username == "teacher01")
        student_id = next(item.id for item in self.app.state.store.users.values() if item.username == "stu001")
        created = self.client.post(
            "/api/admin/courses",
            json={"id": "audit-102", "name": "审计专题", "term": "2026-2027", "teacher_id": teacher_id},
        )
        self.assertEqual(created.status_code, 201)
        self.assertEqual(created.json()["id"], "audit-102")
        updated_course = self.client.put(
            "/api/admin/courses/audit-102",
            json={"name": "审计专题更新", "term": "2026-2027", "teacher_id": teacher_id},
        )
        self.assertEqual(updated_course.status_code, 200)
        self.assertEqual(updated_course.json()["name"], "审计专题更新")
        self.assertEqual(self.client.get("/api/admin/courses").status_code, 200)
        self.assertEqual(
            self.client.delete("/api/admin/courses/audit-101").status_code,
            409,
        )
        self.client.post("/api/auth/logout")

        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "teacher01", "password": "teach123*"}
            ).status_code,
            200,
        )
        self.assertTrue(any(course["id"] == "audit-102" for course in self.client.get("/api/courses").json()))
        enrolled = self.client.post(
            f"/api/courses/audit-102/students/{student_id}", json={"class_name": "审计二班"}
        )
        self.assertEqual(enrolled.status_code, 201)
        self.assertEqual(enrolled.json()["class_name"], "审计二班")
        self.assertEqual(
            self.client.post(
                f"/api/courses/audit-102/students/{student_id}", json={"class_name": "重复"}
            ).status_code,
            409,
        )
        self.assertEqual(len(self.client.get("/api/courses/audit-102/students").json()["students"]), 1)
        self.assertEqual(
            self.client.get("/api/courses/audit-102/available-students").json()["students"],
            [],
        )
        self.client.post("/api/auth/logout")

        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "stu001", "password": "stu123*"}
            ).status_code,
            200,
        )
        self.assertTrue(any(course["id"] == "audit-102" for course in self.client.get("/api/courses").json()))
        self.client.post("/api/auth/logout")

        self.client.post(
            "/api/auth/login", json={"username": "teacher01", "password": "teach123*"}
        )
        self.assertEqual(self.client.delete(f"/api/courses/audit-102/students/{student_id}").status_code, 204)
        self.client.post("/api/auth/logout")
        self.client.post(
            "/api/auth/login", json={"username": "admin", "password": "admin123*"}
        )
        self.assertEqual(self.client.delete("/api/admin/courses/audit-102").status_code, 204)

    def test_admin_can_disable_and_enable_user(self) -> None:
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "admin", "password": "admin123*"}
            ).status_code,
            200,
        )
        student_id = next(item.id for item in self.app.state.store.users.values() if item.username == "stu001")
        disabled = self.client.patch(
            f"/api/admin/users/{student_id}/status", json={"status": "disabled"}
        )
        self.assertEqual(disabled.status_code, 200)
        self.assertEqual(disabled.json()["status"], "disabled")
        self.client.post("/api/auth/logout")
        blocked = self.client.post(
            "/api/auth/login", json={"username": "stu001", "password": "stu123*"}
        )
        self.assertEqual(blocked.status_code, 403)
        self.assertEqual(blocked.json()["code"], "account_disabled")
        self.client.post(
            "/api/auth/login", json={"username": "admin", "password": "admin123*"}
        )
        enabled = self.client.patch(
            f"/api/admin/users/{student_id}/status", json={"status": "active"}
        )
        self.assertEqual(enabled.status_code, 200)
        self.assertEqual(enabled.json()["status"], "active")
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "stu001", "password": "stu123*"}
            ).status_code,
            200,
        )

    def test_user_can_change_password_and_old_password_expires(self) -> None:
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "stu001", "password": "stu123*"}
            ).status_code,
            200,
        )
        changed = self.client.post(
            "/api/auth/change-password",
            json={"current_password": "stu123*", "new_password": "stu45678*"},
        )
        self.assertEqual(changed.status_code, 204)
        self.assertEqual(self.client.get("/api/auth/me").status_code, 401)
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "stu001", "password": "stu123*"}
            ).status_code,
            401,
        )
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "stu001", "password": "stu45678*"}
            ).status_code,
            200,
        )

    def test_admin_backup_and_restore(self) -> None:
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "admin", "password": "admin123*"}
            ).status_code,
            200,
        )
        backup = self.client.get("/api/admin/backup")
        self.assertEqual(backup.status_code, 200)
        self.assertIn("users", backup.json())
        restored = self.client.post(
            "/api/admin/restore",
            files={"file": ("backup.json", backup.content, "application/json")},
        )
        self.assertEqual(restored.status_code, 200)
        self.assertEqual(self.client.get("/api/auth/me").status_code, 401)

    def test_admin_archive_backup_restores_local_files(self) -> None:
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "admin", "password": "admin123*"}
            ).status_code,
            200,
        )
        uploaded = self.client.post(
            "/api/docs/upload",
            data={"target_kb": "audit_textbook"},
            files={"file": ("archive.md", b"# Archive", "text/markdown")},
        )
        self.assertEqual(uploaded.status_code, 202)
        document_id = uploaded.json()["id"]
        archive = self.client.get("/api/admin/backup?format=archive")
        self.assertEqual(archive.status_code, 200)
        self.assertEqual(archive.headers["content-type"], "application/zip")
        self.assertTrue(archive.content.startswith(b"PK"))
        self.assertEqual(self.client.delete(f"/api/docs/{document_id}").status_code, 204)
        self.assertEqual(
            self.client.post(
                "/api/admin/restore",
                files={"file": ("backup.zip", archive.content, "application/zip")},
            ).status_code,
            200,
        )
        self.client.post(
            "/api/auth/login", json={"username": "admin", "password": "admin123*"}
        )
        self.assertTrue(any(item["id"] == document_id for item in self.client.get("/api/docs").json()))

    def test_restore_rejects_invalid_backup(self) -> None:
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "admin", "password": "admin123*"}
            ).status_code,
            200,
        )
        invalid = self.client.post(
            "/api/admin/restore",
            files={"file": ("backup.json", b"[]", "application/json")},
        )
        self.assertEqual(invalid.status_code, 422)
        state = self.app.state.store.export_json()
        payload = json.loads(state)
        payload["documents"] = {"bad": {"file_key": "../../outside.txt"}}
        invalid_path = self.client.post(
            "/api/admin/restore",
            files={"file": ("backup.json", json.dumps(payload).encode(), "application/json")},
        )
        self.assertEqual(invalid_path.status_code, 422)

    def test_json_store_survives_reconstruction(self) -> None:
        with TemporaryDirectory() as data_dir:
            settings = Settings(
                app_env="development",
                data_dir=data_dir,
                persistence_enabled=True,
                cookie_secure=False,
            )
            first = Store(settings)
            first.courses["audit-101"].name = "审计学（已保存）"
            student = first.find_user_by_username("stu001")
            session = QuizSession("assigned-1", student.id, "audit-101", [], assigned=True, assignment_id="batch-1")
            first.quiz_sessions[session.id] = session
            first.assignments["batch-1"] = {"id": "batch-1", "course_id": "audit-101", "session_ids": [session.id]}
            first.roll_calls.append({"course_id": "audit-101", "student_id": student.id, "correct": True})
            first.save()
            self.assertTrue(Path(data_dir, "store.json").exists())
            second = Store(settings)
            self.assertEqual(second.courses["audit-101"].name, "审计学（已保存）")
            self.assertEqual(len(second.questions), 3)
            self.assertEqual(second.quiz_sessions["assigned-1"].assignment_id, "batch-1")
            self.assertIn("batch-1", second.assignments)
            self.assertTrue(second.roll_calls[0]["correct"])
            self.assertEqual(second.sessions, {})

    def test_store_recovers_interrupted_grading_with_submitted_answers(self) -> None:
        with TemporaryDirectory() as data_dir:
            settings = Settings(app_env="development", data_dir=data_dir, persistence_enabled=True, cookie_secure=False)
            first = Store(settings)
            student = first.find_user_by_username("stu001")
            session = QuizSession(
                "interrupted-1", student.id, "audit-101",
                [{"id": "q1", "type": "short_answer", "score": 10, "knowledge_points": []}],
                status="grading", submitted_at=datetime.now(timezone.utc),
                submitted_answers={"q1": "提交后进程中断"},
            )
            first.quiz_sessions[session.id] = session
            self.assertTrue(first.save())

            resumed = Store(settings).quiz_sessions[session.id]

            self.assertEqual(resumed.status, "ongoing")
            self.assertIsNone(resumed.submitted_at)
            self.assertEqual(resumed.submitted_answers, {"q1": "提交后进程中断"})
            persisted = json.loads(Path(data_dir, "store.json").read_text(encoding="utf-8"))
            self.assertEqual(persisted["quiz_sessions"][session.id]["status"], "ongoing")

    def test_llm_provider_config_resolves_key_env_without_repr_leak(self) -> None:
        raw = json.dumps([{
            "name": "student-a", "channel": "student", "base_url": "https://provider.test/v1",
            "api_key_env": "TEST_STUDENT_LLM_KEY", "model": "test-model",
        }])
        with patch.dict(os.environ, {"LLM_PROVIDERS_JSON": raw, "TEST_STUDENT_LLM_KEY": "secret-test-value"}):
            providers = _llm_providers()

        self.assertEqual(len(providers), 1)
        self.assertEqual(providers[0].api_key, "secret-test-value")
        self.assertNotIn("secret-test-value", repr(providers[0]))

    def test_llm_channel_selection_does_not_borrow_teacher_legacy_key(self) -> None:
        student = LLMProvider("student-only", "student", "https://student.test/v1", "student-key", "student-model", 4, 10, 1, "student-only")
        settings = Settings(
            llm_base_url="https://legacy-teacher.test/v1", llm_api_key="legacy-key", llm_model="legacy-model",
            llm_providers=(student,),
        )

        self.assertEqual([item.name for item in LLMClient(settings, "student")._providers()], ["student-only"])
        self.assertEqual(LLMClient(settings, "teacher")._providers(), ())
        self.assertEqual([item.name for item in LLMClient(settings, "grading")._providers()], ["student-only"])
        self.assertFalse(settings.llm_configured_for("teacher"))
        self.assertTrue(settings.llm_configured_for("grading"))

    def test_corrupt_persistent_store_is_not_replaced_with_seed_data(self) -> None:
        with TemporaryDirectory() as data_dir:
            store_path = Path(data_dir, "store.json")
            settings = Settings(app_env="development", data_dir=data_dir, persistence_enabled=True, cookie_secure=False)
            for raw_state in (
                "{broken",
                json.dumps({"version": 2, "users": {}, "courses": {}, "questions": {}, "quiz_sessions": None}),
            ):
                with self.subTest(raw_state=raw_state):
                    store_path.write_text(raw_state, encoding="utf-8")
                    with self.assertRaisesRegex(OSError, "refusing to seed"):
                        Store(settings)
                    self.assertEqual(store_path.read_text(encoding="utf-8"), raw_state)

    def test_concurrent_json_saves_preserve_all_mutations(self) -> None:
        with TemporaryDirectory() as data_dir:
            settings = Settings(app_env="development", data_dir=data_dir, persistence_enabled=True, cookie_secure=False)
            store = Store(settings)

            def write(index):
                store.audit("concurrent-save", None, {"index": index})
                self.assertTrue(store.save())

            with ThreadPoolExecutor(max_workers=8) as pool:
                list(pool.map(write, range(32)))

            loaded = Store(settings)
            saved = {item["detail"]["index"] for item in loaded.audit_logs if item["action"] == "concurrent-save"}
            self.assertEqual(saved, set(range(32)))

    def test_read_requests_do_not_rewrite_persistent_store(self) -> None:
        with TemporaryDirectory() as data_dir:
            settings = Settings(app_env="development", data_dir=data_dir, persistence_enabled=True, cookie_secure=False)
            with patch("app.main.settings", settings):
                with TestClient(create_app()) as client:
                    self.assertEqual(client.get("/health").status_code, 200)
                    self.assertFalse(Path(data_dir, "store.json").exists())
                    self.assertEqual(client.post("/api/auth/login", json={"username": "stu001", "password": "stu123*"}).status_code, 200)
                    self.assertTrue(Path(data_dir, "store.json").exists())

    def test_legacy_store_migration_rebuilds_mastery_and_removes_demo_scores(self) -> None:
        with TemporaryDirectory() as data_dir:
            settings = Settings(app_env="development", data_dir=data_dir, persistence_enabled=True, cookie_secure=False)
            first = Store(settings)
            student = first.find_user_by_username("stu001")
            actual = QuizSession("actual-1", student.id, "audit-101", [{"id": "q1", "score": 10, "knowledge_points": ["审计证据"]}])
            actual.submitted_at = actual.started_at
            demo = QuizSession("demo-1", student.id, "audit-101", [{"id": "q2", "score": 10, "knowledge_points": ["审计证据"]}])
            demo.submitted_at = demo.started_at
            first.quiz_sessions.update({actual.id: actual, demo.id: demo})
            first.results["actual-result"] = GradingResult("actual-result", actual.id, [{"question_id": "q1", "score": 10, "max_score": 10, "method": "rule", "why": "正确", "knowledge_points": ["审计证据"]}], 10, 10, "graded")
            first.results["demo-result"] = GradingResult("demo-result", demo.id, [{"question_id": "q2", "score": 0, "max_score": 10, "method": "rule", "why": "演示", "knowledge_points": ["审计证据"]}], 0, 10, "graded")
            first.mastery[(student.id, "audit-101", "审计证据")] = {"mastery": 30, "attempts": 1}
            first.reports[student.id] = {"markdown": "旧报告"}
            first.audit("quiz_demo_random", "teacher", {"session_id": demo.id})
            state = first._state()
            state["version"] = 1
            Path(data_dir, "store.json").write_text(json.dumps(state), encoding="utf-8")

            migrated = Store(settings)
            self.assertEqual(migrated.mastery[(student.id, "audit-101", "审计证据")]["mastery"], 100)
            self.assertEqual(migrated.mastery[(student.id, "audit-101", "审计证据")]["attempts"], 1)
            self.assertNotIn(demo.id, migrated.quiz_sessions)
            self.assertNotIn("demo-result", migrated.results)
            self.assertFalse(migrated.reports)
            self.assertEqual(json.loads(Path(data_dir, "store.json").read_text(encoding="utf-8"))["version"], 2)

    def test_admin_delete_user_cleans_related_state(self) -> None:
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "stu001", "password": "stu123*"}
            ).status_code,
            200,
        )
        user_id = self.client.get("/api/auth/me").json()["id"]
        self.client.post("/api/chat/ask", json={"question": "待清理聊天"})
        quiz = self.client.post(
            "/api/quiz/start", json={"counts": {"single": 1, "multi": 0, "short": 0}}
        ).json()
        self.client.post(
            f"/api/quiz/{quiz['id']}/submit",
            json={"answers": {quiz["questions"][0]["id"]: 1}},
        )
        self.client.post("/api/auth/logout")
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "admin", "password": "admin123*"}
            ).status_code,
            200,
        )
        self.assertEqual(self.client.delete(f"/api/admin/users/{user_id}").status_code, 204)
        store = self.app.state.store
        self.assertNotIn(user_id, store.users)
        self.assertFalse(any(key[1] == user_id for key in store.enrollments))
        self.assertFalse(any(session.user_id == user_id for session in store.quiz_sessions.values()))
        self.assertFalse(any(item.get("user_id") == user_id for item in store.chat_history))
        self.assertFalse(any(key[0] == user_id for key in store.mastery))

    def test_admin_cannot_delete_teacher_with_courses(self) -> None:
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "admin", "password": "admin123*"}
            ).status_code,
            200,
        )
        teacher_id = next(user.id for user in self.app.state.store.users.values() if user.username == "teacher01")
        response = self.client.delete(f"/api/admin/users/{teacher_id}")
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.json()["code"], "teacher_has_courses")

    def test_manual_grading_review_flow(self) -> None:
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "stu001", "password": "stu123*"}
            ).status_code,
            200,
        )
        quiz = self.client.post(
            "/api/quiz/start",
            json={"course_id": "audit-101", "counts": {"single": 0, "multi": 0, "short": 1}},
        ).json()
        question_id = quiz["questions"][0]["id"]
        result = self.client.post(
            f"/api/quiz/{quiz['id']}/submit",
            json={"answers": {question_id: "充分性是数量要求，适当性是质量要求。"}},
        ).json()["result"]
        self.assertEqual(result["status"], "needs_review")
        self.client.post("/api/auth/logout")

        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "teacher01", "password": "teach123*"}
            ).status_code,
            200,
        )
        self.assertEqual(self.client.get("/api/grading/queue").status_code, 200)
        pending = self.client.get("/api/grading/queue/details").json()
        item = next(row for row in pending if row["result_id"] == result["id"])
        self.assertEqual(item["student"]["display_name"], "张伟")
        self.assertEqual(item["answer"], "充分性是数量要求，适当性是质量要求。")
        self.assertTrue(item["question"]["reference_answer"])
        reviewed = self.client.post(
            f"/api/grading/{result['id']}/review",
            json={"question_id": question_id, "score": 20, "reason": "覆盖了充分性和适当性"},
        )
        self.assertEqual(reviewed.status_code, 200)
        self.assertEqual(reviewed.json()["status"], "graded")

    def test_rubric_suggestion_requires_teacher_review_before_mastery(self) -> None:
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "stu001", "password": "stu123*"}
            ).status_code,
            200,
        )
        quiz = self.client.post(
            "/api/quiz/start",
            json={"course_id": "audit-101", "counts": {"single": 0, "multi": 0, "short": 1}},
        ).json()
        question_id = quiz["questions"][0]["id"]
        submitted = self.client.post(
            f"/api/quiz/{quiz['id']}/submit",
            json={"answers": {question_id: "充分性是数量要求，适当性是质量要求，关系是数量不能弥补质量。"}},
        )
        self.assertEqual(submitted.status_code, 200)
        self.assertEqual(submitted.json()["result"]["status"], "needs_review")
        self.assertEqual(submitted.json()["result"]["per_question"][0]["method"], "manual_pending")
        self.assertNotIn("per_point", submitted.json()["result"]["per_question"][0])
        self.assertFalse(self.app.state.store.mastery)
        self.client.post("/api/auth/logout")
        self.client.post("/api/auth/login", json={"username": "teacher01", "password": "teach123*"})
        reviewed = self.client.post(
            f"/api/grading/{submitted.json()['result']['id']}/review",
            json={"question_id": question_id, "score": 25, "reason": "核对答案覆盖全部评分点"},
        )
        self.assertEqual(reviewed.status_code, 200)
        self.assertTrue(all(len(key) == 3 for key in self.app.state.store.mastery))
        self.client.post("/api/auth/logout")
        self.client.post("/api/auth/login", json={"username": "stu001", "password": "stu123*"})
        report = self.client.post("/api/progress/report")
        self.assertEqual(report.status_code, 200)
        self.assertEqual(report.json()["report"]["source"], "rule")
        self.assertIn("学情报告", report.json()["report"]["markdown"])
        history = self.client.get("/api/progress/history")
        self.assertEqual(history.status_code, 200)
        self.assertEqual(history.json()[0]["status"], "graded")
        self.assertTrue(history.json()[0]["wrong_questions"] == [])

    def test_manual_review_score_limit(self) -> None:
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "stu001", "password": "stu123*"}
            ).status_code,
            200,
        )
        quiz = self.client.post(
            "/api/quiz/start",
            json={"course_id": "audit-101", "counts": {"single": 0, "multi": 0, "short": 1}},
        ).json()
        question_id = quiz["questions"][0]["id"]
        result = self.client.post(
            f"/api/quiz/{quiz['id']}/submit",
            json={"answers": {question_id: "只回答了充分性。"}},
        ).json()["result"]
        self.client.post("/api/auth/logout")
        self.client.post(
            "/api/auth/login", json={"username": "teacher01", "password": "teach123*"}
        )
        self.assertEqual(
            self.client.post(
                f"/api/grading/{result['id']}/review",
                json={"question_id": question_id, "score": 26, "reason": "超过满分"},
            ).status_code,
            422,
        )

    def test_graph_mapping_and_target_cleanup(self) -> None:
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "teacher01", "password": "teach123*"}
            ).status_code,
            200,
        )
        knowledge = self.client.get("/api/graph/nodes?graph=knowledge").json()[0]
        problem = self.client.get("/api/graph/nodes?graph=problem").json()[0]
        mapped = self.client.put(
            "/api/graph/mapping",
            json={"source_id": problem["id"], "target_id": knowledge["id"]},
        )
        self.assertEqual(mapped.status_code, 200)
        self.assertEqual(self.client.delete(f"/api/graph/nodes/{knowledge['id']}").status_code, 204)
        remaining = self.client.get(f"/api/graph/nodes?graph=problem").json()
        self.assertIsNone(next(node for node in remaining if node["id"] == problem["id"])["map_from"])

    def test_graph_mapping_rejects_invalid_targets(self) -> None:
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "teacher01", "password": "teach123*"}
            ).status_code,
            200,
        )
        knowledge = self.client.get("/api/graph/nodes?graph=knowledge").json()[0]
        problem = self.client.get("/api/graph/nodes?graph=problem").json()[0]
        competency = self.client.get("/api/graph/nodes?graph=comp").json()[0]
        self.assertEqual(
            self.client.post(
                "/api/graph/nodes",
                json={"graph": "problem", "branch": "B", "name": "合法问题", "map_from": knowledge["id"]},
            ).status_code,
            201,
        )
        self.assertEqual(
            self.client.post(
                "/api/graph/nodes",
                json={"graph": "problem", "branch": "B", "name": "错误目标", "map_from": competency["id"]},
            ).status_code,
            422,
        )
        self.assertEqual(
            self.client.post(
                "/api/graph/nodes",
                json={"graph": "problem", "branch": "B", "name": "不存在目标", "map_from": "missing-node"},
            ).status_code,
            404,
        )
        self.assertEqual(
            self.client.put(
                f"/api/graph/nodes/{problem['id']}",
                json={"graph": "knowledge", "branch": "A", "name": "非法更新", "map_from": knowledge["id"]},
            ).status_code,
            422,
        )

    def test_graph_mapping_rejects_other_teachers_target(self) -> None:
        store = self.app.state.store
        other = store.add_user("teacher02", "王老师", "teach234*", "teacher")
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "teacher01", "password": "teach123*"}
            ).status_code,
            200,
        )
        target = self.client.post(
            "/api/graph/nodes", json={"graph": "knowledge", "branch": "B", "name": "教师一目标"}
        ).json()
        self.client.post("/api/auth/logout")
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "teacher02", "password": "teach234*"}
            ).status_code,
            200,
        )
        source = self.client.post(
            "/api/graph/nodes", json={"graph": "problem", "branch": "B", "name": "教师二问题"}
        ).json()
        self.assertEqual(
            self.client.put(
                "/api/graph/mapping", json={"source_id": source["id"], "target_id": target["id"]}
            ).status_code,
            403,
        )

    def test_graph_writes_are_owner_scoped(self) -> None:
        store = self.app.state.store
        store.add_user("teacher02", "王老师", "teach234*", "teacher")
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "teacher01", "password": "teach123*"}
            ).status_code,
            200,
        )
        node = self.client.get("/api/graph/nodes?graph=knowledge").json()[0]
        self.client.post("/api/auth/logout")
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "teacher02", "password": "teach234*"}
            ).status_code,
            200,
        )
        self.assertEqual(self.client.get("/api/graph/nodes?graph=knowledge").json(), [])
        self.assertEqual(
            self.client.put(
                f"/api/graph/nodes/{node['id']}",
                json={"graph": "knowledge", "branch": "A", "name": "越权更新"},
            ).status_code,
            403,
        )
        self.assertEqual(self.client.delete(f"/api/graph/nodes/{node['id']}").status_code, 403)

    def test_document_upload_and_delete_uses_local_storage(self) -> None:
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "teacher01", "password": "teach123*"}
            ).status_code,
            200,
        )
        uploaded = self.client.post(
            "/api/docs/upload",
            data={"target_kb": "audit_textbook"},
            files={"file": ("notes.md", b"# Audit notes", "text/markdown")},
        )
        self.assertEqual(uploaded.status_code, 202)
        document_id = uploaded.json()["id"]
        document = self.app.state.store.documents[document_id]
        path = Path(self.app.state.store.settings.data_dir) / document["file_key"]
        self.assertTrue(path.exists())
        self.assertEqual(self.client.delete(f"/api/docs/{document_id}").status_code, 204)
        self.assertFalse(path.exists())

    def test_upload_rejects_mismatched_file_content(self) -> None:
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "teacher01", "password": "teach123*"}
            ).status_code,
            200,
        )
        invalid_pdf = self.client.post(
            "/api/docs/upload",
            data={"target_kb": "audit_textbook"},
            files={"file": ("notes.pdf", b"not a pdf", "application/pdf")},
        )
        self.assertEqual(invalid_pdf.status_code, 422)
        invalid_markdown = self.client.post(
            "/api/docs/upload",
            data={"target_kb": "audit_textbook"},
            files={"file": ("notes.md", b"\xff\xfe", "text/markdown")},
        )
        self.assertEqual(invalid_markdown.status_code, 422)

    def test_teacher_resources_are_owner_scoped(self) -> None:
        store = self.app.state.store
        store.add_user("teacher02", "王老师", "teach234*", "teacher")
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "teacher01", "password": "teach123*"}
            ).status_code,
            200,
        )
        document = self.client.post(
            "/api/docs/upload",
            data={"target_kb": "audit_textbook"},
            files={"file": ("private.md", b"# Private", "text/markdown")},
        ).json()
        self.client.post("/api/auth/logout")

        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "teacher02", "password": "teach234*"}
            ).status_code,
            200,
        )
        self.assertEqual(self.client.get("/api/docs").json(), [])
        self.assertEqual(self.client.delete(f"/api/docs/{document['id']}").status_code, 403)

        self.client.post("/api/auth/logout")
        self.client.post(
            "/api/auth/login", json={"username": "admin", "password": "admin123*"}
        )
        self.assertEqual(self.client.delete(f"/api/docs/{document['id']}").status_code, 204)

    def test_chat_degraded_history_and_usage(self) -> None:
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "stu001", "password": "stu123*"}
            ).status_code,
            200,
        )
        asked = self.client.post("/api/chat/ask", json={"question": "什么是审计证据？"})
        self.assertEqual(asked.status_code, 200)
        self.assertTrue(asked.json()["degraded"])
        self.assertEqual(self.client.get("/api/chat/history?limit=1").status_code, 200)
        self.assertEqual(len(self.client.get("/api/chat/export").json()["items"]), 1)
        self.assertEqual(self.app.state.store.usage_logs[-1]["outcome"], "model_not_configured")
        self.assertNotIn(self.client.get("/api/auth/me").json()["id"], self.app.state.store.chat_quota)
        self.assertEqual(self.client.delete("/api/chat/history").status_code, 204)
        self.assertEqual(self.client.get("/api/chat/history").json(), [])

    def test_chat_course_boundaries_and_monitor_scope(self) -> None:
        store = self.app.state.store
        store.courses["other-course"] = Course("other-course", "其他课程", "2025-2026", "other-teacher")

        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "stu001", "password": "stu123*"}
            ).status_code,
            200,
        )
        self.assertEqual(
            self.client.post(
                "/api/chat/ask", json={"question": "跨课程问题", "course_id": "other-course"}
            ).status_code,
            403,
        )
        self.assertEqual(
            self.client.post(
                "/api/chat/ask", json={"question": "不存在课程", "course_id": "missing-course"}
            ).status_code,
            404,
        )
        asked = self.client.post(
            "/api/chat/ask", json={"question": "课程内问题", "course_id": "audit-101"}
        )
        self.assertEqual(asked.status_code, 200)
        self.assertEqual(self.client.get("/api/chat/history").json()[0]["course_id"], "audit-101")

        self.client.post("/api/auth/logout")
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "teacher01", "password": "teach123*"}
            ).status_code,
            200,
        )
        store.chat_history.append({
            "user_id": next(item.id for item in store.users.values() if item.username == "stu001"),
            "course_id": "other-course",
            "question": "不应出现在该教师监控中的问题",
            "response": {},
        })
        monitored = self.client.get("/api/chat/monitor")
        self.assertEqual(monitored.status_code, 200)
        self.assertTrue(monitored.json())
        self.assertTrue(all(item["course_id"] == "audit-101" for item in monitored.json()))
        self.assertEqual(self.client.get("/api/chat/monitor", params={"course_id": "other-course"}).status_code, 403)
        chat = monitored.json()[0]
        marked = self.client.patch(f"/api/chat/monitor/{chat['id']}/important", json={"important": True})
        self.assertEqual(marked.status_code, 200)
        self.assertTrue(next(item for item in store.chat_history if item["id"] == chat["id"])["important"])

    def test_maxkb_retrieval_mapping(self) -> None:
        class Response:
            def raise_for_status(self) -> None:
                return None

            def json(self) -> dict:
                return {
                    "data": {
                        "records": [{
                            "content": "审计证据应当充分且适当。",
                            "document_name": "审计准则",
                            "similarity": 0.91,
                        }]
                    }
                }

        class Client:
            def __init__(self, **kwargs) -> None:
                self.kwargs = kwargs

            async def __aenter__(self):
                return self

            async def __aexit__(self, *args) -> None:
                return None

            async def post(self, *args, **kwargs) -> Response:
                return Response()

        settings = Settings(
            maxkb_url="http://maxkb.test",
            maxkb_api_key="test-key",
            maxkb_dataset_ids=("audit-textbook",),
        )
        with patch("app.services.maxkb_client.httpx.AsyncClient", Client):
            result = asyncio.run(MaxKBClient(settings).retrieve("什么是审计证据？", "audit-101"))
        self.assertIsNotNone(result)
        self.assertIn("审计证据应当充分且适当", result[0]["text"])
        self.assertEqual(result[0]["name"], "审计准则")

    def test_maxkb_course_mapping_fails_closed(self) -> None:
        settings = Settings(
            maxkb_url="http://maxkb.test",
            maxkb_api_key="test-key",
            maxkb_dataset_ids=("global-dataset",),
            maxkb_course_dataset_ids=(("audit-101", ("course-dataset",)),),
        )
        self.assertIsNone(asyncio.run(MaxKBClient(settings).retrieve("question", "other-course")))
        empty = Settings(
            maxkb_url="http://maxkb.test",
            maxkb_api_key="test-key",
            maxkb_course_dataset_ids=(("audit-101", ()),),
        )
        self.assertIsNone(asyncio.run(MaxKBClient(empty).retrieve("question", "audit-101")))

    def test_question_review_lifecycle_and_snapshot_version(self) -> None:
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "teacher01", "password": "teach123*"}
            ).status_code,
            200,
        )
        created = self.client.post(
            "/api/bank",
            json={
                "type": "single_choice",
                "stem": "导入审核测试题",
                "options": ["A", "B"],
                "answer": 0,
                "reference_answer": "A",
                "knowledge_points": ["测试知识点"],
                "difficulty": "easy",
                "course_id": "audit-101",
            },
        )
        self.assertEqual(created.status_code, 201)
        question_id = created.json()["id"]
        self.assertNotIn("answer", created.json())
        self.assertEqual(self.client.post(f"/api/bank/{question_id}/submit").json()["status"], "reviewing")
        rejected = self.client.post(
            f"/api/bank/{question_id}/review",
            json={"decision": "reject", "comment": "需要补充题干背景"},
        )
        self.assertEqual(rejected.status_code, 200)
        self.assertEqual(rejected.json()["status"], "draft")
        self.assertEqual(rejected.json()["review_comment"], "需要补充题干背景")
        self.assertEqual(self.client.post(f"/api/bank/{question_id}/submit").json()["status"], "reviewing")
        published = self.client.post(
            f"/api/bank/{question_id}/review", json={"decision": "pass"}
        )
        self.assertEqual(published.json()["status"], "published")

        quiz = self.client.post(
            "/api/quiz/start",
            json={"course_id": "audit-101", "counts": {"single": 1, "multi": 0, "short": 0}},
        )
        self.assertEqual(quiz.status_code, 200)
        self.assertEqual(quiz.json()["bank_version"], "audit-101-v1")

    def test_csv_import_requires_confirmation(self) -> None:
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "teacher01", "password": "teach123*"}
            ).status_code,
            200,
        )
        csv_data = (
            "code,type,stem,options,answer,reference_answer,knowledge_points,difficulty\n"
            "IMP-0001,single_choice,CSV导入题,A|B,0,A,审计证据,easy\n"
        )
        preview = self.client.post(
            "/api/bank/import",
            data={"course_id": "audit-101"},
            files={"file": ("questions.csv", csv_data.encode("utf-8"), "text/csv")},
        )
        self.assertEqual(preview.status_code, 200)
        self.assertEqual(preview.json()["valid"], 1)
        self.assertEqual(self.client.get("/api/bank?status=draft").json(), [])
        confirmed = self.client.post(f"/api/bank/import/{preview.json()['preview_id']}/confirm")
        self.assertEqual(confirmed.status_code, 201)
        self.assertEqual(confirmed.json()["count"], 1)
        self.assertEqual(confirmed.json()["questions"][0]["status"], "draft")

    def test_xlsx_import_requires_confirmation(self) -> None:
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "teacher01", "password": "teach123*"}
            ).status_code,
            200,
        )
        workbook = b'''<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Sheet1" sheetId="1" r:id="rId1"/></sheets></workbook>'''
        relationships = b'''<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/></Relationships>'''
        worksheet = b'''<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>code</t></is></c><c r="B1" t="inlineStr"><is><t>type</t></is></c><c r="C1" t="inlineStr"><is><t>stem</t></is></c><c r="D1" t="inlineStr"><is><t>options</t></is></c><c r="E1" t="inlineStr"><is><t>answer</t></is></c><c r="F1" t="inlineStr"><is><t>reference_answer</t></is></c><c r="G1" t="inlineStr"><is><t>knowledge_points</t></is></c><c r="H1" t="inlineStr"><is><t>difficulty</t></is></c></row><row r="2"><c r="A2" t="inlineStr"><is><t>XLSX-0001</t></is></c><c r="B2" t="inlineStr"><is><t>single_choice</t></is></c><c r="C2" t="inlineStr"><is><t>XLSX import test</t></is></c><c r="D2" t="inlineStr"><is><t>A|B</t></is></c><c r="E2"><v>0</v></c><c r="F2" t="inlineStr"><is><t>A</t></is></c><c r="G2" t="inlineStr"><is><t>evidence</t></is></c><c r="H2" t="inlineStr"><is><t>easy</t></is></c></row></sheetData></worksheet>'''
        archive = BytesIO()
        with ZipFile(archive, "w") as xlsx:
            xlsx.writestr("xl/workbook.xml", workbook)
            xlsx.writestr("xl/_rels/workbook.xml.rels", relationships)
            xlsx.writestr("xl/worksheets/sheet1.xml", worksheet)
        preview = self.client.post(
            "/api/bank/import",
            data={"course_id": "audit-101"},
            files={"file": ("questions.xlsx", archive.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        self.assertEqual(preview.status_code, 200)
        self.assertEqual(preview.json()["valid"], 1)
        confirmed = self.client.post(f"/api/bank/import/{preview.json()['preview_id']}/confirm")
        self.assertEqual(confirmed.status_code, 201)
        self.assertEqual(confirmed.json()["questions"][0]["code"], "XLSX-0001")

    def test_invalid_choice_types_are_rejected(self) -> None:
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "teacher01", "password": "teach123*"}
            ).status_code,
            200,
        )
        invalid = self.client.post(
            "/api/bank",
            json={
                "type": "single_choice",
                "stem": "错误答案类型",
                "options": ["A", "B"],
                "answer": True,
                "reference_answer": "A",
            },
        )
        self.assertEqual(invalid.status_code, 422)

    def test_llm_unconfigured_keeps_rule_grading_and_degraded_chat(self) -> None:
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "stu001", "password": "stu123*"}
            ).status_code,
            200,
        )
        quiz = self.client.post(
            "/api/quiz/start",
            json={"course_id": "audit-101", "counts": {"single": 0, "multi": 0, "short": 1}},
        ).json()
        question_id = quiz["questions"][0]["id"]
        submitted = self.client.post(
            f"/api/quiz/{quiz['id']}/submit",
            json={"answers": {question_id: "充分性是数量要求，适当性是质量要求，关系是数量不能弥补质量。"}},
        )
        self.assertEqual(submitted.status_code, 200)
        methods = {item["method"] for item in submitted.json()["result"]["per_question"]}
        self.assertTrue(methods.issubset({"rubric_rule", "manual_pending"}))
        asked = self.client.post("/api/chat/ask", json={"question": "什么是审计证据？"})
        self.assertEqual(asked.status_code, 200)
        self.assertTrue(asked.json()["degraded"])

    def test_chat_retries_unsupported_json_mode_and_reads_text_parts(self) -> None:
        self.app.state.store.settings = Settings(
            llm_base_url="https://relay.example/v1", llm_api_key="test-key", llm_model="test-model"
        )
        self.assertEqual(
            self.client.post("/api/auth/login", json={"username": "stu001", "password": "stu123*"}).status_code,
            200,
        )
        requests = []
        answer = json.dumps({
            "answer_markdown": "审计证据需要充分且适当。",
            "sections": {"conclusion": "审计证据需要充分且适当。"},
            "mind_map": ["审计证据", "└─ 充分且适当"],
        }, ensure_ascii=False)

        def relay(request):
            requests.append(request)
            payload = json.loads(request.content)
            if "response_format" in payload:
                return httpx.Response(
                    400, json={"error": {"message": "response_format json_object is not supported"}}, request=request
                )
            middle = len(answer) // 2
            return httpx.Response(200, json={"choices": [{"message": {"content": [
                {"type": "text", "text": answer[:middle]}, {"type": "text", "text": answer[middle:]},
            ]}}], "usage": {"prompt_tokens": 19, "completion_tokens": 7}}, headers={"x-request-id": "relay-test-1"}, request=request)

        transport = httpx.MockTransport(relay)
        original_client = httpx.AsyncClient
        with patch(
            "app.services.llm_client.httpx.AsyncClient",
            side_effect=lambda **kwargs: original_client(transport=transport, **kwargs),
        ):
            response = self.client.post("/api/chat/ask", json={"question": "什么是审计证据？", "course_id": "audit-101"})

        self.assertEqual(response.status_code, 200)
        self.assertIn("充分且适当", response.json()["answer_markdown"])
        self.assertNotEqual(response.json()["degraded_reason"], "model_unconfigured")
        self.assertEqual(len(requests), 2)
        self.assertEqual(str(requests[0].url), "https://relay.example/v1/chat/completions")
        self.assertIn("response_format", json.loads(requests[0].content))
        self.assertNotIn("response_format", json.loads(requests[1].content))
        usage = self.app.state.store.usage_logs[-1]
        self.assertEqual(usage["model"], "test-model")
        self.assertEqual((usage["prompt_tokens"], usage["completion_tokens"]), (19, 7))

    def test_chat_distinguishes_invalid_gateway_response_from_unconfigured(self) -> None:
        self.app.state.store.settings = Settings(
            llm_base_url="https://relay.example/v1", llm_api_key="test-key", llm_model="test-model"
        )
        self.client.post("/api/auth/login", json={"username": "stu001", "password": "stu123*"})
        transport = httpx.MockTransport(
            lambda request: httpx.Response(
                200, json={"data": {"answer": "有响应但不是兼容格式"}},
                headers={"x-request-id": "relay-test-invalid"}, request=request,
            )
        )
        original_client = httpx.AsyncClient
        with self.assertLogs("app.services.llm_client", level="WARNING") as logs:
            with patch(
                "app.services.llm_client.httpx.AsyncClient",
                side_effect=lambda **kwargs: original_client(transport=transport, **kwargs),
            ):
                response = self.client.post("/api/chat/ask", json={"question": "什么是审计证据？", "course_id": "audit-101"})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["degraded_reason"], "model_invalid_response")
        self.assertIn("已配置", response.json()["answer_markdown"])
        self.assertTrue(any("relay-test-invalid" in message for message in logs.output))

    def test_chat_exposes_relay_rate_limit(self) -> None:
        self.app.state.store.settings = Settings(
            llm_base_url="https://relay-rate.example/v1", llm_api_key="test-key", llm_model="test-model"
        )
        self.client.post("/api/auth/login", json={"username": "stu001", "password": "stu123*"})
        transport = httpx.MockTransport(lambda request: httpx.Response(
            429, json={"error": {"message": "rate limited"}},
            headers={"retry-after": "3", "x-request-id": "relay-test-429"}, request=request,
        ))
        original_client = httpx.AsyncClient
        with self.assertLogs("app.services.llm_client", level="WARNING") as logs:
            with patch(
                "app.services.llm_client.httpx.AsyncClient",
                side_effect=lambda **kwargs: original_client(transport=transport, **kwargs),
            ):
                response = self.client.post("/api/chat/ask", json={"question": "什么是审计证据？", "course_id": "audit-101"})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["degraded_reason"], "model_rate_limited")
        self.assertTrue(any("retry_after=3" in message and "relay-test-429" in message for message in logs.output))
        user_id = self.client.get("/api/auth/me").json()["id"]
        self.assertEqual(self.app.state.store.chat_quota[user_id]["count"], 1)

    def test_sync_llm_client_retries_unsupported_json_mode(self) -> None:
        client = LLMClient(Settings(
            llm_base_url="https://relay.example/v1", llm_api_key="test-key", llm_model="test-model"
        ))
        calls = []

        def relay(request):
            calls.append(request)
            payload = json.loads(request.content)
            if "response_format" in payload:
                return httpx.Response(
                    422, json={"error": {"message": "response_format is unsupported"}}, request=request
                )
            return httpx.Response(
                200, json={"choices": [{"message": {"content": "{\"ok\":true}"}}]}, request=request
            )

        original_client = httpx.Client
        with patch(
            "app.services.llm_client.httpx.Client",
            side_effect=lambda **kwargs: original_client(transport=httpx.MockTransport(relay), **kwargs),
        ):
            result = client.complete_json([{"role": "user", "content": "return JSON"}])

        self.assertEqual(result, {"ok": True})
        self.assertEqual(len(calls), 2)

    def test_llm_fails_over_from_rate_limited_student_provider(self) -> None:
        providers = (
            LLMProvider("pool-primary", "student", "https://primary.test/v1", "primary-key", "model-a", 1, 4, 10, "pool-primary"),
            LLMProvider("pool-backup", "student", "https://backup.test/v1", "backup-key", "model-b", 1, 4, 20, "pool-backup"),
        )
        settings = Settings(llm_base_url="", llm_api_key="", llm_model="", llm_providers=providers)
        client = LLMClient(settings, channel="student")
        calls = []

        def relay(request):
            calls.append(request.url.host)
            if request.url.host == "primary.test":
                return httpx.Response(429, json={"error": {"message": "limit"}}, request=request)
            return httpx.Response(200, json={
                "choices": [{"message": {"content": "{\"ok\":true}"}}],
                "usage": {"prompt_tokens": 9, "completion_tokens": 4},
            }, request=request)

        original_client = httpx.Client
        with patch(
            "app.services.llm_client.httpx.Client",
            side_effect=lambda **kwargs: original_client(transport=httpx.MockTransport(relay), **kwargs),
        ):
            result = client.complete_json([{"role": "user", "content": "return JSON"}])

        self.assertEqual(result, {"ok": True})
        self.assertEqual(calls, ["primary.test", "backup.test"])
        self.assertEqual(client.last_provider, "pool-backup")
        self.assertEqual(client.last_model, "model-b")
        self.assertEqual(client.last_usage, {"prompt_tokens": 9, "completion_tokens": 4})

    def test_llm_provider_gate_caps_shared_sync_requests(self) -> None:
        provider = LLMProvider(
            "gate-test", "shared", "https://gate.test/v1", "test-key", "test-model",
            3, 16, 1, "gate-test-unique",
        )
        settings = Settings(llm_base_url="", llm_api_key="", llm_model="", llm_providers=(provider,))
        active = 0
        peak = 0
        guard = Lock()

        def relay(request):
            nonlocal active, peak
            with guard:
                active += 1
                peak = max(peak, active)
            sleep(0.02)
            with guard:
                active -= 1
            return httpx.Response(200, json={"choices": [{"message": {"content": "ok"}}]}, request=request)

        original_client = httpx.Client
        with patch(
            "app.services.llm_client.httpx.Client",
            side_effect=lambda **kwargs: original_client(transport=httpx.MockTransport(relay), **kwargs),
        ):
            with ThreadPoolExecutor(max_workers=12) as pool:
                responses = list(pool.map(
                    lambda _: LLMClient(settings).complete([{"role": "user", "content": "hello"}]),
                    range(12),
                ))

        self.assertEqual(responses, ["ok"] * 12)
        self.assertEqual(peak, 3)

    def test_transient_timeout_does_not_discard_queued_provider_request(self) -> None:
        provider = LLMProvider(
            "timeout-queue-test", "shared", "https://timeout-queue.test/v1", "test-key", "test-model",
            1, 2, 1, "timeout-queue-test-unique",
        )
        settings = Settings(llm_base_url="", llm_api_key="", llm_model="", llm_providers=(provider,))
        calls = 0
        guard = Lock()

        def relay(request):
            nonlocal calls
            with guard:
                calls += 1
                call_number = calls
            if call_number == 1:
                sleep(0.02)
                raise httpx.ReadTimeout("transient", request=request)
            return httpx.Response(200, json={"choices": [{"message": {"content": "ok"}}]}, request=request)

        original_client = httpx.Client
        with patch(
            "app.services.llm_client.httpx.Client",
            side_effect=lambda **kwargs: original_client(transport=httpx.MockTransport(relay), **kwargs),
        ):
            with ThreadPoolExecutor(max_workers=2) as pool:
                results = list(pool.map(
                    lambda _: LLMClient(settings).complete([{"role": "user", "content": "hello"}]),
                    range(2),
                ))

        self.assertEqual(calls, 2)
        self.assertIn("ok", results)
        self.assertIn(None, results)

    def test_llm_provider_gate_caps_async_requests(self) -> None:
        provider = LLMProvider(
            "async-gate-test", "shared", "https://async-gate.test/v1", "test-key", "test-model",
            3, 16, 1, "async-gate-test-unique",
        )
        settings = Settings(llm_base_url="", llm_api_key="", llm_model="", llm_providers=(provider,))

        async def run_requests():
            class Tracker:
                def __init__(self):
                    self.active = 0
                    self.peak = 0
                    self.lock = asyncio.Lock()

            tracker = Tracker()

            class Transport(httpx.AsyncBaseTransport):
                async def handle_async_request(self, request):
                    async with tracker.lock:
                        tracker.active += 1
                        tracker.peak = max(tracker.peak, tracker.active)
                    await asyncio.sleep(0.02)
                    async with tracker.lock:
                        tracker.active -= 1
                    return httpx.Response(
                        200, json={"choices": [{"message": {"content": "ok"}}]}, request=request
                    )

            original_client = httpx.AsyncClient
            with patch(
                "app.services.llm_client.httpx.AsyncClient",
                side_effect=lambda **kwargs: original_client(transport=Transport(), **kwargs),
            ):
                results = await asyncio.gather(*(
                    LLMClient(settings).complete_async([{"role": "user", "content": "hello"}])
                    for _ in range(12)
                ))
            return results, tracker.peak

        results, peak = asyncio.run(run_requests())
        self.assertEqual(results, ["ok"] * 12)
        self.assertEqual(peak, 3)

    def test_async_llm_fails_over_to_backup(self) -> None:
        providers = (
            LLMProvider("async-primary", "student", "https://async-primary.test/v1", "key-a", "model-a", 1, 2, 1, "async-primary"),
            LLMProvider("async-backup", "student", "https://async-backup.test/v1", "key-b", "model-b", 1, 2, 2, "async-backup"),
        )
        settings = Settings(llm_base_url="", llm_api_key="", llm_model="", llm_providers=providers)
        client = LLMClient(settings, channel="student")
        calls = []

        def relay(request):
            calls.append(request.url.host)
            if request.url.host == "async-primary.test":
                return httpx.Response(503, json={"error": "maintenance"}, request=request)
            return httpx.Response(200, json={"choices": [{"message": {"content": "{\"ready\":true}"}}]}, request=request)

        transport = httpx.MockTransport(relay)
        original_client = httpx.AsyncClient
        with patch(
            "app.services.llm_client.httpx.AsyncClient",
            side_effect=lambda **kwargs: original_client(transport=transport, **kwargs),
        ):
            result = asyncio.run(client.complete_json_async([{"role": "user", "content": "return JSON"}]))

        self.assertEqual(result, {"ready": True})
        self.assertEqual(calls, ["async-primary.test", "async-backup.test"])

    def test_llm_provider_attempt_count_is_bounded(self) -> None:
        providers = tuple(
            LLMProvider(f"attempt-{index}", "shared", f"https://attempt-{index}.test/v1", "key", "model", 1, 2, index, f"attempt-{index}")
            for index in range(4)
        )
        settings = Settings(
            llm_base_url="", llm_api_key="", llm_model="", llm_providers=providers,
            llm_max_attempts=2, llm_total_timeout_seconds=5,
        )
        calls = []

        def relay(request):
            calls.append(request.url.host)
            return httpx.Response(503, json={"error": "maintenance"}, request=request)

        original_client = httpx.Client
        with patch(
            "app.services.llm_client.httpx.Client",
            side_effect=lambda **kwargs: original_client(transport=httpx.MockTransport(relay), **kwargs),
        ):
            client = LLMClient(settings, "shared")
            result = client.complete([{"role": "user", "content": "hello"}])

        self.assertIsNone(result)
        self.assertEqual(calls, ["attempt-0.test", "attempt-1.test"])
        self.assertEqual(client.last_error, "provider_error")

    def test_async_llm_enforces_total_request_deadline(self) -> None:
        provider = LLMProvider(
            "deadline-test", "shared", "https://deadline.test/v1", "key", "model", 1, 2, 1, "deadline-test",
        )
        settings = Settings(
            llm_base_url="", llm_api_key="", llm_model="", llm_providers=(provider,),
            llm_timeout_seconds=5, llm_total_timeout_seconds=0.05, llm_queue_timeout_seconds=1,
        )

        class SlowTransport(httpx.AsyncBaseTransport):
            async def handle_async_request(self, request):
                await asyncio.sleep(0.2)
                return httpx.Response(200, json={"choices": [{"message": {"content": "too late"}}]}, request=request)

        original_client = httpx.AsyncClient
        with patch(
            "app.services.llm_client.httpx.AsyncClient",
            side_effect=lambda **kwargs: original_client(transport=SlowTransport(), **kwargs),
        ):
            client = LLMClient(settings, "shared")
            started = time.perf_counter()
            result = asyncio.run(client.complete_async([{"role": "user", "content": "hello"}]))
            elapsed = time.perf_counter() - started

        self.assertIsNone(result)
        self.assertEqual(client.last_error, "timeout")
        self.assertLess(elapsed, 0.5)

    def test_llm_rate_limits_absent_by_default(self) -> None:
        provider = LLMProvider(
            "no-quota-test", "shared", "https://no-quota.test/v1", "key", "model", 4, 8, 1, "no-quota-test-unique",
        )
        self.assertIsNone(_quota_for(provider))
        limited = LLMProvider(
            "with-quota-test", "shared", "https://with-quota.test/v1", "key", "model", 4, 8, 1,
            "with-quota-test-unique", rpm_limit=10, tpm_limit=1000,
        )
        self.assertIsNotNone(_quota_for(limited))

    def test_llm_rpm_limit_throttles_requests(self) -> None:
        provider = LLMProvider(
            "rpm-test", "shared", "https://rpm.test/v1", "key", "model", 8, 16, 1, "rpm-test-unique", rpm_limit=2,
        )
        settings = Settings(
            llm_base_url="", llm_api_key="", llm_model="", llm_providers=(provider,),
            llm_queue_timeout_seconds=0.2, llm_total_timeout_seconds=5,
        )
        calls = 0
        guard = Lock()

        def relay(request):
            nonlocal calls
            with guard:
                calls += 1
            return httpx.Response(200, json={"choices": [{"message": {"content": "ok"}}]}, request=request)

        def run(_: int) -> tuple[str | None, str | None]:
            client = LLMClient(settings)
            return client.complete([{"role": "user", "content": "hello"}]), client.last_error

        original_client = httpx.Client
        with patch(
            "app.services.llm_client.httpx.Client",
            side_effect=lambda **kwargs: original_client(transport=httpx.MockTransport(relay), **kwargs),
        ):
            with ThreadPoolExecutor(max_workers=6) as pool:
                outcomes = list(pool.map(run, range(6)))

        self.assertEqual(calls, 2)
        self.assertEqual([result for result, _ in outcomes].count("ok"), 2)
        self.assertEqual([error for _, error in outcomes if error], ["rate_limited"] * 4)

    def test_llm_tpm_limit_reconciles_with_actual_usage(self) -> None:
        provider = LLMProvider(
            "tpm-test", "shared", "https://tpm.test/v1", "key", "model", 4, 8, 1, "tpm-test-unique", tpm_limit=100,
        )
        settings = Settings(
            llm_base_url="", llm_api_key="", llm_model="", llm_providers=(provider,),
            llm_token_output_estimate=50, llm_queue_timeout_seconds=0.2, llm_total_timeout_seconds=5,
        )
        calls = 0

        def relay(request):
            nonlocal calls
            calls += 1
            return httpx.Response(200, json={
                "choices": [{"message": {"content": "ok"}}],
                "usage": {"prompt_tokens": 5, "completion_tokens": 5},
            }, request=request)

        original_client = httpx.Client
        with patch(
            "app.services.llm_client.httpx.Client",
            side_effect=lambda **kwargs: original_client(transport=httpx.MockTransport(relay), **kwargs),
        ):
            client = LLMClient(settings)
            # Estimate per request is 66 tokens of the 100 budget; each response bills only
            # 10, so the returned surplus keeps the next request admissible. Without the
            # reconciliation the second request would already be throttled.
            results = [client.complete([{"role": "user", "content": "hello"}]) for _ in range(5)]

        self.assertEqual(results, ["ok"] * 4 + [None])
        self.assertEqual(calls, 4)
        self.assertEqual(client.last_error, "rate_limited")

    def test_llm_rate_limited_provider_falls_over_to_backup(self) -> None:
        providers = (
            LLMProvider("quota-primary", "student", "https://qprimary.test/v1", "key-a", "model-a", 2, 4, 10, "quota-primary-unique", rpm_limit=1),
            LLMProvider("quota-backup", "student", "https://qbackup.test/v1", "key-b", "model-b", 2, 4, 20, "quota-backup-unique"),
        )
        settings = Settings(
            llm_base_url="", llm_api_key="", llm_model="", llm_providers=providers,
            llm_queue_timeout_seconds=0.2, llm_total_timeout_seconds=5,
        )
        calls = []

        def relay(request):
            calls.append(request.url.host)
            return httpx.Response(200, json={"choices": [{"message": {"content": "ok"}}]}, request=request)

        original_client = httpx.Client
        with patch(
            "app.services.llm_client.httpx.Client",
            side_effect=lambda **kwargs: original_client(transport=httpx.MockTransport(relay), **kwargs),
        ):
            first = LLMClient(settings, "student")
            self.assertEqual(first.complete([{"role": "user", "content": "hello"}]), "ok")
            self.assertEqual(first.last_provider, "quota-primary")
            second = LLMClient(settings, "student")
            self.assertEqual(second.complete([{"role": "user", "content": "hello"}]), "ok")
            self.assertEqual(second.last_provider, "quota-backup")

        self.assertEqual(calls, ["qprimary.test", "qbackup.test"])

    def test_llm_quota_refund_after_unbilled_attempt(self) -> None:
        provider = LLMProvider(
            "refund-test", "shared", "https://refund.test/v1", "key", "model", 4, 8, 1, "refund-test-unique", tpm_limit=100,
        )
        settings = Settings(
            llm_base_url="", llm_api_key="", llm_model="", llm_providers=(provider,),
            llm_token_output_estimate=50, llm_queue_timeout_seconds=0.2, llm_total_timeout_seconds=5,
        )
        calls = 0
        guard = Lock()

        def relay(request):
            nonlocal calls
            with guard:
                calls += 1
                current = calls
            if current == 1:
                raise httpx.ReadTimeout("transient", request=request)
            return httpx.Response(200, json={"choices": [{"message": {"content": "ok"}}]}, request=request)

        original_client = httpx.Client
        with patch(
            "app.services.llm_client.httpx.Client",
            side_effect=lambda **kwargs: original_client(transport=httpx.MockTransport(relay), **kwargs),
        ):
            client = LLMClient(settings)
            self.assertIsNone(client.complete([{"role": "user", "content": "hello"}]))
            self.assertEqual(client.complete([{"role": "user", "content": "hello"}]), "ok")

        self.assertEqual(calls, 2)

    def test_async_llm_rpm_limit_throttles_requests(self) -> None:
        provider = LLMProvider(
            "async-rpm-test", "shared", "https://async-rpm.test/v1", "key", "model", 8, 16, 1, "async-rpm-test-unique", rpm_limit=1,
        )
        settings = Settings(
            llm_base_url="", llm_api_key="", llm_model="", llm_providers=(provider,),
            llm_queue_timeout_seconds=0.2, llm_total_timeout_seconds=5,
        )

        async def run() -> list[str | None]:
            transport = httpx.MockTransport(
                lambda request: httpx.Response(200, json={"choices": [{"message": {"content": "ok"}}]}, request=request)
            )
            original_client = httpx.AsyncClient
            with patch(
                "app.services.llm_client.httpx.AsyncClient",
                side_effect=lambda **kwargs: original_client(transport=transport, **kwargs),
            ):
                return list(await asyncio.gather(*(
                    LLMClient(settings).complete_async([{"role": "user", "content": "hello"}])
                    for _ in range(4)
                )))

        results = asyncio.run(run())
        self.assertEqual(results.count("ok"), 1)
        self.assertEqual(results.count(None), 3)

    def test_demo_random_grades_without_changing_student_records(self) -> None:
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "teacher01", "password": "teach123*"}
            ).status_code,
            200,
        )
        demo = self.client.post(
            "/api/quiz/demo-random",
            json={"course_id": "audit-101", "counts": {"single": 1, "multi": 1, "short": 1}},
        )
        self.assertEqual(demo.status_code, 200)
        body = demo.json()
        self.assertEqual(body["student"]["display_name"], "张伟")
        self.assertEqual(len(body["questions"]), 3)
        self.assertIn(body["result"]["status"], {"graded", "needs_review"})
        store = self.app.state.store
        self.assertTrue(body["demo"])
        self.assertNotIn(body["session_id"], store.quiz_sessions)
        self.assertFalse(store.mastery)
        self.assertFalse(store.results)

    def test_first_mastery_score_uses_actual_result(self) -> None:
        self.client.post("/api/auth/login", json={"username": "stu001", "password": "stu123*"})
        question = next(item for item in self.app.state.store.questions.values() if item.type == "single_choice")
        quiz = self.client.post(
            "/api/quiz/start",
            json={"course_id": "audit-101", "counts": {"single": 1, "multi": 0, "short": 0}},
        ).json()
        question_id = quiz["questions"][0]["id"]
        submitted = self.client.post(f"/api/quiz/{quiz['id']}/submit", json={"answers": {question_id: question.answer}})
        self.assertEqual(submitted.status_code, 200)
        mastery = next(iter(self.app.state.store.mastery.values()))
        self.assertEqual(mastery["mastery"], 100)
        self.assertEqual(mastery["attempts"], 1)

    def test_teacher_can_search_students_and_open_individual_progress(self) -> None:
        self.client.post("/api/auth/login", json={"username": "teacher01", "password": "teach123*"})
        directory = self.client.get("/api/courses/audit-101/students", params={"q": "张伟", "page_size": 10})
        self.assertEqual(directory.status_code, 200)
        self.assertEqual(directory.json()["total"], 1)
        student = directory.json()["students"][0]
        self.assertEqual(student["display_name"], "张伟")
        progress = self.client.get(f"/api/progress/students/{student['id']}", params={"course_id": "audit-101"})
        self.assertEqual(progress.status_code, 200)
        self.assertEqual(progress.json()["student"]["id"], student["id"])
        self.assertEqual(self.client.get(f"/api/progress/students/{student['id']}", params={"course_id": "missing"}).status_code, 404)

    def test_classroom_response_is_saved_to_student_record(self) -> None:
        self.client.post("/api/auth/login", json={"username": "teacher01", "password": "teach123*"})
        student_id = self.app.state.store.students_for_course("audit-101")[0].id
        saved = self.client.post("/api/courses/audit-101/roll-call", json={"student_id": student_id, "question": "什么是审计证据？", "correct": True})
        self.assertEqual(saved.status_code, 201)
        self.assertEqual(self.client.get("/api/courses/audit-101/roll-call").json()[0]["student_id"], student_id)
        progress = self.client.get(f"/api/progress/students/{student_id}", params={"course_id": "audit-101"})
        self.assertEqual(progress.json()["classroom"][0]["question"], "什么是审计证据？")

    def test_excel_archive_schedule_and_safe_history_cleanup(self) -> None:
        with TemporaryDirectory() as data_dir:
            settings = Settings(app_env="development", data_dir=data_dir, persistence_enabled=True, cookie_secure=False)
            with patch("app.main.settings", settings):
                app = create_app()
                with TestClient(app) as client:
                    self.assertEqual(client.post("/api/auth/login", json={"username": "teacher01", "password": "teach123*"}).status_code, 200)
                    snapshot = client.post("/api/admin/excel/snapshot", json={})
                    self.assertEqual(snapshot.status_code, 200)
                    self.assertIn("spreadsheetml.sheet", snapshot.headers["content-type"])
                    workbook = load_workbook(BytesIO(snapshot.content), read_only=True)
                    self.assertTrue({"概览", "学生名单", "作业布置", "测验记录", "答题明细", "掌握度", "答疑记录", "课堂互动", "AI用量", "操作日志"}.issubset(workbook.sheetnames))
                    self.assertEqual(client.put("/api/admin/excel/schedule", json={"enabled": True, "interval_hours": 24}).status_code, 200)
                    teacher = app.state.store.find_user_by_username("teacher01")
                    app.state.store.excel_schedules[teacher.id]["next_at"] = (datetime.now(timezone.utc) - timedelta(seconds=1)).isoformat()

                    class StopAfterSnapshot:
                        calls = 0

                        def wait(self, _seconds):
                            self.calls += 1
                            return self.calls > 1

                    run_snapshot_scheduler(app.state.store, StopAfterSnapshot())
                    self.assertEqual(len(app.state.store.excel_snapshots), 2)
                    self.assertGreater(datetime.fromisoformat(app.state.store.excel_schedules[teacher.id]["next_at"]), datetime.now(timezone.utc))

                    client.post("/api/auth/logout")
                    client.post("/api/auth/login", json={"username": "stu001", "password": "stu123*"})
                    quiz = client.post("/api/quiz/start", json={"course_id": "audit-101", "counts": {"single": 1, "multi": 0, "short": 0}}).json()
                    question_id = quiz["questions"][0]["id"]
                    self.assertEqual(client.post(f"/api/quiz/{quiz['id']}/submit", json={"answers": {question_id: 1}}).status_code, 200)
                    session = app.state.store.quiz_sessions[quiz["id"]]
                    session.submitted_at = datetime.now(timezone.utc) - timedelta(days=30)
                    session.started_at = session.submitted_at - timedelta(minutes=10)
                    client.post("/api/auth/logout")
                    client.post("/api/auth/login", json={"username": "teacher01", "password": "teach123*"})
                    before = (datetime.now(timezone.utc) - timedelta(days=1)).isoformat()
                    query = {"before": before, "categories": ["quiz"]}
                    preview = client.get("/api/admin/history/cleanup-preview", params=query)
                    self.assertEqual(preview.status_code, 200)
                    self.assertEqual(preview.json()["counts"]["quiz"], 1)
                    cleaned = client.post("/api/admin/history/cleanup", json={**query, "confirmed": True})
                    self.assertEqual(cleaned.status_code, 200)
                    self.assertNotIn(quiz["id"], app.state.store.quiz_sessions)
                    self.assertIsNotNone(cleaned.json()["snapshot_id"])
                    self.assertEqual(client.get("/api/admin/excel/snapshots").status_code, 200)

    def test_quick_question_requires_teacher_and_returns_answer(self) -> None:
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "stu001", "password": "stu123*"}
            ).status_code,
            200,
        )
        self.assertEqual(self.client.get("/api/bank/quick-question").status_code, 403)
        self.client.post("/api/auth/logout")
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "teacher01", "password": "teach123*"}
            ).status_code,
            200,
        )
        question = self.client.get("/api/bank/quick-question", params={"course_id": "audit-101"})
        self.assertEqual(question.status_code, 200)
        self.assertTrue(question.json()["answer_display"])
        self.assertTrue(question.json()["stem"])

    def test_ai_question_generation_creates_drafts_from_selected_points(self) -> None:
        self.assertEqual(
            self.client.post(
                "/api/auth/login", json={"username": "teacher01", "password": "teach123*", "role": "teacher"}
            ).status_code,
            200,
        )

        class FakeLLM:
            configured = True

            def __init__(self, settings, channel=None):
                self.settings = settings

            def complete_json(self, messages, temperature=0.2):
                return {
                    "questions": [{
                        "stem": "下列哪项属于可靠的审计证据？",
                        "options": ["银行询证函回函", "未经核实的传闻"],
                        "answer": 0,
                        "reference_answer": "银行询证函回函",
                        "rubric": [],
                    }]
                }

        class FakeKnowledge:
            def search(self, query, top_k=8):
                return [{"name": "审计准则.md", "text": "银行询证函回函属于外部审计证据。", "score": 1.0}]

        self.app.state.store.knowledge = FakeKnowledge()
        with patch("app.routers.bank.LLMClient", FakeLLM):
            generated = self.client.post(
                "/api/bank/generate",
                json={
                    "course_id": "audit-101",
                    "knowledge_points": ["审计证据可靠性"],
                    "question_type": "single_choice",
                    "difficulty": "easy",
                    "count": 1,
                },
            )
        self.assertEqual(generated.status_code, 201)
        body = generated.json()
        self.assertEqual(body["status"], "draft")
        self.assertEqual(len(body["questions"]), 1)
        self.assertEqual(body["questions"][0]["status"], "draft")
        self.assertNotIn("answer", body["questions"][0])
        self.assertEqual(body["questions"][0]["knowledge_points"], ["审计证据可靠性"])

    def test_local_knowledge_base_build_and_search(self) -> None:
        with TemporaryDirectory() as tmp:
            source = Path(tmp) / "kb"
            (source / "电子版教材PDF").mkdir(parents=True)
            (source / "电子版教材PDF" / "教材.md").write_text(
                "审计证据是注册会计师为了得出审计结论而使用的所有信息，函证是获取外部证据的重要程序。",
                encoding="utf-8",
            )
            (source / "审计准则汇总").mkdir(parents=True)
            (source / "审计准则汇总" / "1312号.md").write_text(
                "函证准则要求注册会计师对询证函的设计、发出和回收保持控制。",
                encoding="utf-8",
            )
            index_path = Path(tmp) / "index.pkl"
            result = build_index(source, index_path, suffixes={".md"})
            self.assertEqual(result["files"], 2)
            self.assertGreater(result["chunks"], 0)

            kb = LocalKnowledgeBase(Settings(knowledge_index_path=str(index_path)))
            self.assertTrue(kb.load())
            hits = kb.search("函证程序如何控制", top_k=2)
            self.assertTrue(hits)
            self.assertIn("函证", hits[0]["text"])
            self.assertTrue(hits[0]["name"].endswith(".md"))
            self.assertEqual(kb.stats()["chunks"], result["chunks"])

    def test_chat_uses_local_knowledge_base_sources(self) -> None:
        class FakeClient:
            def __init__(self, *args, **kwargs) -> None:
                pass

            @property
            def configured(self) -> bool:
                return True

            async def complete_json_async(self, *args, **kwargs) -> dict:
                return {
                    "answer_markdown": "函证需要保持控制。",
                    "sections": {"conclusion": "函证需要保持控制。"},
                    "mind_map": ["函证", "└─ 控制"],
                }

        with TemporaryDirectory() as tmp:
            source = Path(tmp) / "kb"
            source.mkdir()
            (source / "准则.md").write_text(
                "函证准则：注册会计师应当对函证过程保持控制，确保回函直接寄回。", encoding="utf-8"
            )
            index_path = Path(tmp) / "index.pkl"
            build_index(source, index_path, suffixes={".md"})
            self.app.state.store.knowledge = LocalKnowledgeBase(
                Settings(knowledge_index_path=str(index_path))
            )

            self.assertEqual(
                self.client.post(
                    "/api/auth/login", json={"username": "stu001", "password": "stu123*"}
                ).status_code,
                200,
            )
            with patch("app.services.chat_service.LLMClient", FakeClient):
                asked = self.client.post(
                    "/api/chat/ask", json={"question": "函证如何保持控制？", "course_id": "audit-101"}
                )
            self.assertEqual(asked.status_code, 200)
            body = asked.json()
            self.assertFalse(body["degraded"])
            self.assertEqual(body["sources"][0]["name"], "准则.md")
            self.assertIn("函证", body["answer_markdown"])

    def test_login_role_must_match_account(self) -> None:
        mismatch = self.client.post(
            "/api/auth/login",
            json={"username": "stu001", "password": "stu123*", "role": "teacher"},
        )
        self.assertEqual(mismatch.status_code, 403)
        self.assertEqual(mismatch.json()["code"], "role_mismatch")
        matched = self.client.post(
            "/api/auth/login",
            json={"username": "stu001", "password": "stu123*", "role": "student"},
        )
        self.assertEqual(matched.status_code, 200)
        self.assertEqual(matched.json()["user"]["role"], "student")

    def test_register_requires_approval_and_prevents_duplicates(self) -> None:
        register = self.client.post(
            "/api/auth/register",
            json={
                "username": "stu2026", "password": "Passw0rd!", "phone": "13800001111",
                "class_code": "AUD101", "display_name": "新同学", "student_number": "STU-2026-001",
            },
        )
        self.assertEqual(register.status_code, 202)
        self.assertEqual(register.json()["status"], "pending")
        pending_user = next(item for item in self.app.state.store.users.values() if item.username == "stu2026")
        self.assertEqual(pending_user.student_number, "STU-2026-001")

        # 待审核账号不能登录，且提示明确
        blocked = self.client.post(
            "/api/auth/login",
            json={"username": "stu2026", "password": "Passw0rd!", "role": "student"},
        )
        self.assertEqual(blocked.status_code, 403)
        self.assertEqual(blocked.json()["code"], "pending_approval")

        # 重复账号 / 重复手机号 / 错误课堂代码
        self.assertEqual(
            self.client.post(
                "/api/auth/register",
                json={"username": "stu2026", "password": "Passw0rd!", "phone": "13800002222", "class_code": "AUD101"},
            ).status_code,
            409,
        )
        self.assertEqual(
            self.client.post(
                "/api/auth/register",
                json={"username": "stu2027", "password": "Passw0rd!", "phone": "13800001111", "class_code": "AUD101"},
            ).status_code,
            409,
        )
        self.assertEqual(
            self.client.post(
                "/api/auth/register",
                json={"username": "stu2029", "password": "Passw0rd!", "phone": "13800002229", "class_code": "AUD101", "student_number": "STU-2026-001"},
            ).status_code,
            409,
        )
        self.assertEqual(
            self.client.post(
                "/api/auth/register",
                json={"username": "stu2028", "password": "Passw0rd!", "phone": "13800003333", "class_code": "WRONG1"},
            ).status_code,
            404,
        )

        # 学生看不到待审核列表
        self.client.post("/api/auth/login", json={"username": "stu001", "password": "stu123*"})
        self.assertEqual(self.client.get("/api/enrollment/pending").status_code, 403)
        self.client.post("/api/auth/logout")

        # 教师可审核，通过后学生可登录
        self.client.post("/api/auth/login", json={"username": "teacher01", "password": "teach123*"})
        pending = self.client.get("/api/enrollment/pending")
        self.assertEqual(pending.status_code, 200)
        row = next(item for item in pending.json() if item["username"] == "stu2026")
        self.assertEqual(row["requested_course_name"], "审计学")
        approved = self.client.post(
            f"/api/enrollment/{row['id']}/approve", json={"class_name": "审计二班"}
        )
        self.assertEqual(approved.status_code, 200)
        self.assertEqual(approved.json()["status"], "active")
        students = self.client.get("/api/courses/audit-101/students").json()["students"]
        self.assertIn("stu2026", [item["username"] for item in students])
        self.client.post("/api/auth/logout")

        login = self.client.post(
            "/api/auth/login",
            json={"username": "stu2026", "password": "Passw0rd!", "role": "student"},
        )
        self.assertEqual(login.status_code, 200)
        me = self.client.get("/api/auth/me").json()
        self.assertEqual([course["id"] for course in me["courses"]], ["audit-101"])

    def test_teacher_can_reject_registration(self) -> None:
        self.assertEqual(
            self.client.post(
                "/api/auth/register",
                json={"username": "stu9999", "password": "Passw0rd!", "phone": "13900009999", "class_code": "AUD101"},
            ).status_code,
            202,
        )
        self.client.post("/api/auth/login", json={"username": "teacher01", "password": "teach123*"})
        row = next(item for item in self.client.get("/api/enrollment/pending").json() if item["username"] == "stu9999")
        rejected = self.client.post(f"/api/enrollment/{row['id']}/reject", json={"reason": "非本班学生"})
        self.assertEqual(rejected.status_code, 204)
        self.assertEqual(
            self.client.get("/api/enrollment/pending").json(),
            [],
        )
        # 被驳回后账号可重新注册
        self.assertEqual(
            self.client.post(
                "/api/auth/register",
                json={"username": "stu9999", "password": "Passw0rd!", "phone": "13900009999", "class_code": "AUD101"},
            ).status_code,
            202,
        )

    def test_registration_rate_limit_blocks_bursts(self) -> None:
        store = self.app.state.store
        store.registration_attempts.clear()
        statuses = []
        for index in range(store.settings.registration_rate_limit + 2):
            statuses.append(
                self.client.post(
                    "/api/auth/register",
                    json={
                        "username": f"burst{index:02d}", "password": "Passw0rd!",
                        "phone": f"1370000{index:04d}", "class_code": "AUD101",
                    },
                ).status_code
            )
        self.assertEqual(statuses[-1], 429)
        self.assertIn(202, statuses)
        store.registration_attempts.clear()

    def test_class_code_rotation_invalidates_old_code(self) -> None:
        self.client.post("/api/auth/login", json={"username": "teacher01", "password": "teach123*"})
        rotated = self.client.post("/api/courses/audit-101/class-code")
        self.assertEqual(rotated.status_code, 200)
        new_code = rotated.json()["class_code"]
        self.assertNotEqual(new_code, "AUD101")
        self.assertEqual(
            self.client.post(
                "/api/auth/register",
                json={"username": "oldcode", "password": "Passw0rd!", "phone": "13600001111", "class_code": "AUD101"},
            ).status_code,
            404,
        )
        self.assertEqual(
            self.client.post(
                "/api/auth/register",
                json={"username": "newcode", "password": "Passw0rd!", "phone": "13600002222", "class_code": new_code},
            ).status_code,
            202,
        )



    # ===== AI 练习(practice) =====
    PRACTICE_FAKE_JSON = {
        "questions": [
            {"type": "single_choice", "stem": "下列哪项属于可靠的审计证据？",
             "options": ["银行询证函回函", "口头答复", "内部记账凭证", "未经核实的传闻"], "answer": 0,
             "reference_answer": "银行询证函回函经银行确认,可靠性最高", "knowledge_points": ["审计证据可靠性"]},
            {"type": "judge", "stem": "审计证据越多越好。", "answer": False,
             "reference_answer": "证据需要适当与充分,并非越多越好", "knowledge_points": ["审计证据适当性"]},
            {"type": "fill", "stem": "函证是获取______证据的重要程序。", "answer": "外部",
             "reference_answer": "函证面向外部第三方获取书面证据", "knowledge_points": ["函证"]},
        ]
    }

    class _PracticeFakeLLM:
        configured = True

        def __init__(self, settings, channel=None):
            self.settings = settings
            self.last_error = None
            self.last_model = "fake-model"
            self.last_request_sent = True
            self.last_usage = {"prompt_tokens": 10, "completion_tokens": 5}

        async def complete_json_async(self, messages, temperature=0.2):
            return json.loads(json.dumps(SmokeTests.PRACTICE_FAKE_JSON))

    class _PracticeBrokenLLM(_PracticeFakeLLM):
        def __init__(self, settings, channel=None):
            super().__init__(settings, channel)
            self.last_error = "rate_limited"
            self.last_request_sent = True

        async def complete_json_async(self, messages, temperature=0.2):
            return None

    def _seed_practice_history(self, count: int = 3) -> None:
        store = self.app.state.store
        student = store.find_user_by_username("stu001")
        for index in range(count):
            store.chat_history.append({
                "id": f"chat-seed-{index}",
                "created_at": datetime.now(timezone.utc).isoformat(),
                "important": False,
                "user_id": student.id,
                "course_id": "audit-101",
                "question": f"审计证据相关问题{index}",
                "response": {"answer_markdown": "x", "sections": {}, "mind_map": [], "sources": [], "degraded": False, "degraded_reason": None},
            })

    def test_practice_generate_answer_and_history_flow(self) -> None:
        self.client.post("/api/auth/login", json={"username": "stu001", "password": "stu123*", "role": "student"})
        self._seed_practice_history()
        with patch("app.routers.practice.LLMClient", self._PracticeFakeLLM):
            generated = self.client.post("/api/practice/generate", json={"course_id": "audit-101", "count": 3})
        self.assertEqual(generated.status_code, 200)
        body = generated.json()
        self.assertEqual(len(body["questions"]), 3)
        self.assertEqual(body["status"], "ongoing")
        for question in body["questions"]:
            self.assertNotIn("answer", question)
        practice_id = body["id"]

        listed = self.client.get("/api/practice").json()
        self.assertEqual([item["id"] for item in listed], [practice_id])
        self.assertEqual(listed[0]["question_count"], 3)
        self.assertEqual(listed[0]["status"], "ongoing")

        first_question = body["questions"][0]
        wrong = self.client.post(f"/api/practice/{practice_id}/submit", json={"answers": {first_question["id"]: "乱写"}})
        self.assertEqual(wrong.status_code, 200)
        self.assertEqual(wrong.json()["result"]["total"], 0)
        self.assertEqual(wrong.json()["status"], "graded")
        duplicate = self.client.post(f"/api/practice/{practice_id}/submit", json={"answers": {}})
        self.assertEqual(duplicate.status_code, 409)

        with patch("app.routers.practice.LLMClient", self._PracticeFakeLLM):
            second = self.client.post("/api/practice/generate", json={"course_id": "audit-101", "count": 3})
        second_id = second.json()["id"]
        stored = self.app.state.store.practices[second_id]
        correct_answers = {question["id"]: question["answer"] for question in stored.questions}
        right = self.client.post(f"/api/practice/{second_id}/submit", json={"answers": correct_answers})
        self.assertEqual(right.status_code, 200)
        self.assertEqual(right.json()["result"]["total"], right.json()["result"]["max_total"])
        self.assertEqual(self.client.get(f"/api/practice/{second_id}").json()["status"], "graded")

        unknown = self.client.post("/api/practice/fake-id/submit", json={"answers": {}})
        self.assertEqual(unknown.status_code, 404)

    def test_practice_generate_requires_chat_history(self) -> None:
        self.client.post("/api/auth/login", json={"username": "stu001", "password": "stu123*", "role": "student"})
        with patch("app.routers.practice.LLMClient", self._PracticeFakeLLM):
            response = self.client.post("/api/practice/generate", json={"count": 3})
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.json()["code"], "no_chat_history")

    def test_practice_generate_model_failure_returns_503(self) -> None:
        self.client.post("/api/auth/login", json={"username": "stu001", "password": "stu123*", "role": "student"})
        self._seed_practice_history()
        with patch("app.routers.practice.LLMClient", self._PracticeBrokenLLM):
            response = self.client.post("/api/practice/generate", json={"count": 3})
        self.assertEqual(response.status_code, 503)
        self.assertEqual(response.json()["code"], "model_rate_limited")
        with patch("app.routers.practice.LLMClient", self._PracticeFakeLLM):
            retry = self.client.post("/api/practice/generate", json={"count": 3})
        self.assertEqual(retry.status_code, 200)

    def test_practice_generate_consumes_chat_quota(self) -> None:
        self.client.post("/api/auth/login", json={"username": "stu001", "password": "stu123*", "role": "student"})
        self._seed_practice_history()
        store = self.app.state.store
        object.__setattr__(store.settings, "chat_daily_limit", 2)
        with patch("app.routers.practice.LLMClient", self._PracticeFakeLLM):
            self.assertEqual(self.client.post("/api/practice/generate", json={"count": 3}).status_code, 200)
            self.assertEqual(self.client.post("/api/practice/generate", json={"count": 3}).status_code, 200)
            third = self.client.post("/api/practice/generate", json={"count": 3})
        self.assertEqual(third.status_code, 429)
        self.assertEqual(third.json()["code"], "chat_quota_exceeded")

    def test_practice_ownership_is_enforced(self) -> None:
        self.client.post("/api/auth/login", json={"username": "stu001", "password": "stu123*", "role": "student"})
        self._seed_practice_history()
        with patch("app.routers.practice.LLMClient", self._PracticeFakeLLM):
            generated = self.client.post("/api/practice/generate", json={"course_id": "audit-101", "count": 3})
        practice_id = generated.json()["id"]

        rival_username = f"stu-rival-{int(time.time()) % 100000}"
        admin_client = TestClient(self.app)
        admin_client.post("/api/auth/login", json={"username": "admin", "password": "admin123*", "role": "admin"})
        admin_client.post("/api/admin/users", json={
            "username": rival_username, "display_name": "对手学生", "password": "rival-pass*1",
            "role": "student", "student_number": f"R{int(time.time()) % 100000}",
        })
        rival_client = TestClient(self.app)
        rival_client.post("/api/auth/login", json={"username": rival_username, "password": "rival-pass*1", "role": "student"})

        self.assertEqual(rival_client.get(f"/api/practice/{practice_id}").status_code, 404)
        self.assertEqual(rival_client.post(f"/api/practice/{practice_id}/submit", json={"answers": {}}).status_code, 404)
        self.assertEqual(rival_client.delete(f"/api/practice/{practice_id}").status_code, 404)
        self.assertEqual(rival_client.get("/api/practice").json(), [])

    def test_practice_delete_and_clear(self) -> None:
        self.client.post("/api/auth/login", json={"username": "stu001", "password": "stu123*", "role": "student"})
        self._seed_practice_history()
        with patch("app.routers.practice.LLMClient", self._PracticeFakeLLM):
            first = self.client.post("/api/practice/generate", json={"count": 3}).json()
            second = self.client.post("/api/practice/generate", json={"count": 3}).json()
        self.assertEqual(len(self.client.get("/api/practice").json()), 2)
        self.assertEqual(self.client.delete(f"/api/practice/{first['id']}").status_code, 204)
        self.assertEqual([item["id"] for item in self.client.get("/api/practice").json()], [second["id"]])
        cleared = self.client.delete("/api/practice")
        self.assertEqual(cleared.status_code, 200)
        self.assertEqual(cleared.json()["deleted"], 1)
        self.assertEqual(self.client.get("/api/practice").json(), [])

    def test_logout_purges_practice_history(self) -> None:
        self.client.post("/api/auth/login", json={"username": "stu001", "password": "stu123*", "role": "student"})
        self._seed_practice_history()
        with patch("app.routers.practice.LLMClient", self._PracticeFakeLLM):
            self.client.post("/api/practice/generate", json={"count": 3})
        self.client.post("/api/auth/logout")
        self.client.post("/api/auth/login", json={"username": "stu001", "password": "stu123*", "role": "student"})
        self.assertEqual(self.client.get("/api/practice").json(), [])
        self.assertTrue(any(item["action"] == "practice_purge_on_logout" for item in self.app.state.store.audit_logs))

    def test_practice_generate_from_single_source_question(self) -> None:
        self.client.post("/api/auth/login", json={"username": "stu001", "password": "stu123*", "role": "student"})
        with patch("app.routers.practice.LLMClient", self._PracticeFakeLLM):
            response = self.client.post("/api/practice/generate", json={
                "count": 2, "source_question": "什么是审计证据",
            })
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(body["source_questions"], ["什么是审计证据"])
        self.assertEqual(body["status"], "ongoing")
        self.assertEqual(len(body["questions"]), 2)

    def test_teacher_cannot_generate_practice(self) -> None:
        self.client.post("/api/auth/login", json={"username": "teacher01", "password": "teach123*", "role": "teacher"})
        response = self.client.post("/api/practice/generate", json={"count": 3})
        self.assertEqual(response.status_code, 403)


if __name__ == "__main__":
    unittest.main()
