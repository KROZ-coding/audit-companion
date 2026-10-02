from copy import deepcopy
from datetime import datetime, timedelta, timezone
from io import BytesIO
import json
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from ..models import User
from ..store import Store


def workbook_bytes(store: Store, user: User, course_id: str | None = None) -> tuple[bytes, list[str]]:
    with store.lock:
        courses = store.courses_for_user(user)
        allowed = {item.id for item in courses}
        if course_id is not None:
            if course_id not in store.courses:
                raise ValueError("课程不存在")
            if user.role != "admin" and course_id not in allowed:
                raise PermissionError("无权导出该课程")
            allowed = {course_id}
            courses = [store.courses[course_id]]
        if user.role == "admin" and course_id is None:
            courses = list(store.courses.values())
            allowed = set(store.courses)
        course_names = {item.id: item.name for item in courses}
        enrollments = deepcopy({key: value for key, value in store.enrollments.items() if key[0] in allowed})
        student_ids = {student_id for _, student_id in enrollments}
        users = deepcopy(store.users)
        sessions = deepcopy([item for item in store.quiz_sessions.values() if item.course_id in allowed])
        session_ids = {item.id for item in sessions}
        results = deepcopy([item for item in store.results.values() if item.session_id in session_ids])
        assignments = deepcopy([item for item in store.assignments.values() if item.get("course_id") in allowed])
        chats = deepcopy([item for item in store.chat_history if item.get("course_id") in allowed])
        roll_calls = deepcopy([item for item in store.roll_calls if item.get("course_id") in allowed])
        mastery = deepcopy([item for item in store.mastery.items() if item[0][1] in allowed or item[0][1] is None])
        questions = deepcopy([item for item in store.questions.values() if item.course_id is None or item.course_id in allowed])
        documents = deepcopy([
            item for item in store.documents.values()
            if user.role == "admin" or item.get("uploaded_by") == user.id or item.get("shared")
        ])
        usage_logs = deepcopy([
            item for item in store.usage_logs
            if user.role == "admin"
            or item.get("user_id") == user.id
            or item.get("course_id") in allowed
        ])
        visible_session_ids = session_ids
        audit_logs = deepcopy([
            item for item in store.audit_logs
            if user.role == "admin"
            or item.get("actor_id") == user.id
            or item.get("detail", {}).get("course_id") in allowed
            or item.get("detail", {}).get("student_id") in student_ids
            or item.get("detail", {}).get("session_id") in visible_session_ids
        ])
        pending = deepcopy([
            item for item in store.users.values()
            if item.status == "pending" and item.requested_course_id in allowed
        ])

    result_by_session = {item.session_id: item for item in results}
    users_by_id = users
    assignments_by_id = {item["id"]: item for item in assignments}
    workbook = Workbook()
    workbook.remove(workbook.active)
    counts: dict[str, int] = {}

    def sheet(title: str, headers: list[str], rows: list[list[object]]) -> None:
        ws = workbook.create_sheet(title[:31])
        ws.append(headers)
        for row in rows:
            ws.append([_cell_value(value) for value in row])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(vertical="center")
        for index, header in enumerate(headers, start=1):
            ws.column_dimensions[ws.cell(1, index).column_letter].width = min(36, max(12, len(header) * 2 + 2))
        counts[title] = len(rows)

    session_rows = []
    answer_rows = []
    for session in sorted(sessions, key=lambda item: item.started_at):
        result = result_by_session.get(session.id)
        student = users_by_id.get(session.user_id)
        assignment = assignments_by_id.get(session.assignment_id or "")
        session_rows.append([
            session.id, student.display_name if student else "学生记录已删除",
            student.student_number if student else "", student.username if student else "",
            session.course_id, course_names.get(session.course_id, ""),
            enrollments.get((session.course_id, session.user_id), ""), session.title,
            "已完成" if result and result.status == "graded" else "待复核" if result else "处理中" if session.status == "grading" else "待作答",
            result.total if result else None, result.max_total if result else None,
            result.graded_by if result else "", session.started_at.isoformat(),
            session.submitted_at.isoformat() if session.submitted_at else "",
            assignment.get("id", "") if assignment else "",
            assignment.get("due_at", "") if assignment else "",
        ])
        if not result:
            continue
        questions_by_id = {item.get("id"): item for item in session.questions}
        for grade in result.per_question:
            question = questions_by_id.get(grade.get("question_id"), {})
            answer_rows.append([
                session.id, session.title, student.display_name if student else "学生记录已删除",
                student.student_number if student else "", session.course_id, course_names.get(session.course_id, ""),
                question.get("id", ""), question.get("stem", ""), question.get("type", ""),
                result.submitted_answers.get(grade.get("question_id")), question.get("reference_answer", ""),
                grade.get("score"), grade.get("max_score"), grade.get("method"), grade.get("why"),
                question.get("knowledge_points", []), question.get("rubric", []),
            ])

    assignment_rows = []
    for assignment in assignments:
        for session_id in assignment.get("session_ids", []):
            session = next((item for item in sessions if item.id == session_id), None)
            if session is None:
                assignment_rows.append([
                    assignment.get("id"), assignment.get("title"), assignment.get("course_id"),
                    course_names.get(assignment.get("course_id"), ""), "", "", "", "记录已清理", None, None,
                    assignment.get("created_at"), assignment.get("due_at"), assignment.get("question_count"),
                ])
                continue
            result = result_by_session.get(session.id)
            student = users_by_id.get(session.user_id)
            assignment_rows.append([
                assignment.get("id"), assignment.get("title"), assignment.get("course_id"),
                course_names.get(assignment.get("course_id"), ""),
                student.display_name if student else "学生记录已删除",
                student.student_number if student else "",
                enrollments.get((session.course_id, session.user_id), ""),
                "已完成" if result and result.status == "graded" else "待复核" if result else "待作答",
                result.total if result else None, result.max_total if result else None,
                assignment.get("created_at"), assignment.get("due_at"), assignment.get("question_count"),
            ])

    sheet("课程", ["课程编号", "课程名称", "学期", "负责人"], [
        [item.id, item.name, item.term, users_by_id.get(item.teacher_id).display_name if users_by_id.get(item.teacher_id) else item.teacher_id]
        for item in courses
    ])
    sheet("学生名单", ["姓名", "学号", "账号", "角色状态", "课程", "课程名称", "班级", "最近登录"], [
        [student.display_name, student.student_number, student.username, student.status,
         cid, course_names.get(cid, ""), class_name, student.last_login_at.isoformat() if student.last_login_at else ""]
        for (cid, student_id), class_name in enrollments.items()
        if (student := users_by_id.get(student_id)) is not None
    ])
    sheet("待审核注册", ["姓名", "学号", "账号", "手机号", "课程", "班级", "申请时间"], [
        [item.display_name, item.student_number, item.username, item.phone, item.requested_course_id,
         course_names.get(item.requested_course_id, ""), item.requested_at.isoformat() if item.requested_at else ""]
        for item in pending
    ])
    sheet("作业布置", ["作业编号", "测验名称", "课程编号", "课程", "学生", "学号", "班级", "状态", "得分", "满分", "布置时间", "截止时间", "题数"], assignment_rows)
    sheet("测验记录", ["会话编号", "学生", "学号", "账号", "课程编号", "课程", "班级", "测验", "状态", "得分", "满分", "批改方式", "开始时间", "提交时间", "作业编号", "截止时间"], session_rows)
    sheet("答题明细", ["会话编号", "测验", "学生", "学号", "课程编号", "课程", "题目编号", "题干", "题型", "学生答案", "参考答案", "得分", "满分", "评分方式", "评分说明", "知识点", "Rubric"], answer_rows)
    mastery_rows = []
    for (student_id, cid, point), value in mastery:
        student = users_by_id.get(student_id)
        mastery_rows.append([
            student.display_name if student else "学生记录已删除", student.student_number if student else "",
            cid, course_names.get(cid, ""), point, value.get("mastery"), value.get("attempts"), value.get("last_updated"),
        ])
    sheet("掌握度", ["学生", "学号", "课程编号", "课程", "知识点", "掌握度", "作答次数", "更新时间"], mastery_rows)
    sheet("答疑记录", ["时间", "课程编号", "课程", "学生", "学号", "问题", "助教回答", "来源", "重点", "降级回答"], [
        [item.get("created_at", ""), item.get("course_id", ""), course_names.get(item.get("course_id"), ""),
         (users_by_id.get(item.get("user_id")) or User("", "", "学生记录已删除", "", "student")).display_name,
         (users_by_id.get(item.get("user_id")) or User("", "", "", "", "student")).student_number,
         item.get("question", ""), (item.get("response") or {}).get("answer_markdown", ""),
         (item.get("response") or {}).get("sources", []), bool(item.get("important")), bool((item.get("response") or {}).get("degraded"))]
        for item in chats
    ])
    sheet("课堂互动", ["时间", "课程编号", "课程", "学生", "学号", "班级", "问题", "结果", "教师"], [
        [item.get("created_at"), item.get("course_id"), course_names.get(item.get("course_id"), ""),
         item.get("student_name"), item.get("student_number"), item.get("class_name"), item.get("question"),
         "答对" if item.get("correct") else "未答对", users_by_id.get(item.get("teacher_id")).display_name if users_by_id.get(item.get("teacher_id")) else item.get("teacher_id")]
        for item in roll_calls
    ])
    sheet("题库", ["编号", "课程", "题型", "题干", "选项", "答案", "参考答案", "Rubric", "知识点", "难度", "状态", "章节", "版本", "来源"], [
        [item.code, course_names.get(item.course_id, "共享题库"), item.type, item.stem, item.options,
         item.answer, item.reference_answer, item.rubric, item.knowledge_points, item.difficulty,
         item.status, item.chapter, item.version, item.source]
        for item in questions
    ])
    sheet("资料库", ["资料名", "知识库", "大小（字节）", "状态", "上传时间", "来源路径", "摘要", "主题"], [
        [item.get("name"), item.get("target_kb"), item.get("size"), item.get("vector_status"),
         item.get("uploaded_at"), item.get("source_path"), item.get("summary"), item.get("topics")]
        for item in documents
    ])
    sheet("AI用量", ["时间", "用户编号", "用户", "课程编号", "用途", "模型", "提示词 token", "输出 token", "耗时毫秒", "结果"], [
        [item.get("at"), item.get("user_id"), users_by_id.get(item.get("user_id")).display_name if users_by_id.get(item.get("user_id")) else "",
         item.get("course_id"), item.get("intent"), item.get("model"), item.get("prompt_tokens"),
         item.get("completion_tokens"), item.get("latency_ms"), item.get("outcome")]
        for item in usage_logs
    ])
    sheet("操作日志", ["时间", "操作人编号", "操作人", "动作", "业务详情"], [
        [item.get("at"), item.get("actor_id"), users_by_id.get(item.get("actor_id")).display_name if users_by_id.get(item.get("actor_id")) else "",
         item.get("action"), item.get("detail", {})]
        for item in audit_logs
    ])
    sheet("概览", ["项目", "记录数或值"], [
        ["导出时间（UTC）", datetime.now(timezone.utc).isoformat()],
        ["导出人", user.display_name], ["导出范围", "全部可访问课程" if course_id is None else course_names.get(course_id, course_id)],
        *[[name, count] for name, count in counts.items()],
    ])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue(), sorted(allowed)


def save_snapshot(store: Store, user: User, course_id: str | None = None) -> tuple[dict, bytes]:
    content, course_ids = workbook_bytes(store, user, course_id)
    snapshot_id = str(uuid4())
    file_key = f"excel_snapshots/{snapshot_id}.xlsx"
    path = Path(store.settings.data_dir) / file_key
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".tmp")
    temporary.write_bytes(content)
    temporary.replace(path)
    now = datetime.now(timezone.utc)
    metadata = {
        "id": snapshot_id, "owner_id": user.id, "file_key": file_key,
        "filename": f"audit_records_{now.strftime('%Y%m%d_%H%M%S')}.xlsx",
        "created_at": now.isoformat(), "bytes": len(content), "course_ids": course_ids,
    }
    remove_files = []
    save_failed = False
    with store._save_lock:
        with store.lock:
            previous_schedule = deepcopy(store.excel_schedules.get(user.id))
            store.excel_snapshots.append(metadata)
            owned = [item for item in store.excel_snapshots if item.get("owner_id") == user.id]
            expired = owned[:-20]
            schedule = store.excel_schedules.get(user.id)
            if schedule and schedule.get("enabled"):
                interval = int(schedule.get("interval_hours", 24))
                schedule["last_saved_at"] = now.isoformat()
                schedule["next_at"] = (now.replace(microsecond=0) + timedelta(hours=interval)).isoformat()
            store.audit("excel_snapshot", user.id, {"snapshot_id": snapshot_id, "course_ids": course_ids})
        if not store.save():
            with store.lock:
                store.excel_snapshots[:] = [item for item in store.excel_snapshots if item.get("id") != snapshot_id]
                if previous_schedule is not None:
                    store.excel_schedules[user.id] = previous_schedule
                store.audit_logs[:] = [item for item in store.audit_logs if item.get("detail", {}).get("snapshot_id") != snapshot_id]
            save_failed = True
        elif expired:
            expired_ids = {item["id"] for item in expired}
            with store.lock:
                store.excel_snapshots[:] = [item for item in store.excel_snapshots if item.get("id") not in expired_ids]
            if store.save():
                remove_files = [Path(store.settings.data_dir) / item["file_key"] for item in expired]
            else:
                with store.lock:
                    existing_ids = {item.get("id") for item in store.excel_snapshots}
                    store.excel_snapshots.extend(item for item in expired if item.get("id") not in existing_ids)
    if save_failed:
        path.unlink(missing_ok=True)
        raise OSError("snapshot metadata could not be persisted")
    for old_path in remove_files:
        old_path.unlink(missing_ok=True)
    return metadata, content


def snapshot_path(store: Store, metadata: dict) -> Path:
    relative = metadata.get("file_key", "")
    if not store._safe_relative_path(relative):
        raise ValueError("invalid snapshot path")
    return Path(store.settings.data_dir) / relative


def cleanup_counts(store: Store, course_ids: set[str], before: datetime, categories: set[str]) -> dict[str, int]:
    with store.lock:
        sessions_to_remove, assignments_to_remove = _completed_quiz_history(store, course_ids, before)
        return {
            "quiz": len(sessions_to_remove) if "quiz" in categories else 0,
            "chat": sum(
                "chat" in categories and item.get("course_id") in course_ids and _before(item.get("created_at"), before)
                for item in store.chat_history
            ),
            "classroom": sum(
                "classroom" in categories and item.get("course_id") in course_ids and _before(item.get("created_at"), before)
                for item in store.roll_calls
            ),
            "usage": sum(
                "usage" in categories and item.get("course_id") in course_ids and _before(item.get("at"), before)
                for item in store.usage_logs
            ),
        }


def cleanup_history(store: Store, user: User, course_ids: set[str], before: datetime, categories: set[str]) -> dict[str, int]:
    with store.lock:
        sessions_to_remove, assignments_to_remove = _completed_quiz_history(store, course_ids, before)
        if "quiz" not in categories:
            sessions_to_remove.clear()
            assignments_to_remove.clear()
        chat_to_remove = {
            id(item) for item in store.chat_history
            if "chat" in categories and item.get("course_id") in course_ids and _before(item.get("created_at"), before)
        }
        classroom_to_remove = {
            id(item) for item in store.roll_calls
            if "classroom" in categories and item.get("course_id") in course_ids and _before(item.get("created_at"), before)
        }
        usage_to_remove = {
            id(item) for item in store.usage_logs
            if "usage" in categories and item.get("course_id") in course_ids and _before(item.get("at"), before)
        }
        counts = {"quiz": len(sessions_to_remove), "chat": len(chat_to_remove), "classroom": len(classroom_to_remove), "usage": len(usage_to_remove)}
        result_ids = {key for key, result in store.results.items() if result.session_id in sessions_to_remove}
        for session_id in sessions_to_remove:
            store.quiz_sessions.pop(session_id, None)
        for result_id in result_ids:
            store.results.pop(result_id, None)
        for assignment_id in assignments_to_remove:
            store.assignments.pop(assignment_id, None)
        store.chat_history[:] = [item for item in store.chat_history if id(item) not in chat_to_remove]
        store.roll_calls[:] = [item for item in store.roll_calls if id(item) not in classroom_to_remove]
        store.usage_logs[:] = [item for item in store.usage_logs if id(item) not in usage_to_remove]
        if sessions_to_remove:
            store.rebuild_mastery()
        store.audit("history_cleanup", user.id, {
            "course_ids": sorted(course_ids), "before": before.isoformat(), "counts": counts,
        })
        return counts


def run_snapshot_scheduler(store: Store, stop_event) -> None:
    while not stop_event.wait(30):
        now = datetime.now(timezone.utc)
        with store.lock:
            due = [
                owner_id for owner_id, schedule in store.excel_schedules.items()
                if schedule.get("enabled") and schedule.get("next_at")
                and datetime.fromisoformat(schedule["next_at"]) <= now
            ]
            users = {owner_id: store.users.get(owner_id) for owner_id in due}
        for owner_id in due:
            user = users.get(owner_id)
            if user is None:
                continue
            try:
                save_snapshot(store, user)
            except Exception as error:
                with store.lock:
                    schedule = store.excel_schedules.get(owner_id)
                    if schedule:
                        schedule["next_at"] = (now + timedelta(hours=1)).isoformat()
                    store.audit("excel_auto_snapshot_failed", owner_id, {"error": str(error)[:300]})
                store.save()


def _completed_quiz_history(store: Store, course_ids: set[str], before: datetime) -> tuple[set[str], set[str]]:
    results_by_session = {item.session_id: item for item in store.results.values()}
    sessions_by_id = store.quiz_sessions
    session_ids: set[str] = set()
    assignment_ids: set[str] = set()
    for assignment_id, assignment in store.assignments.items():
        if assignment.get("course_id") not in course_ids or not _before(assignment.get("created_at"), before):
            continue
        assigned = [sessions_by_id.get(session_id) for session_id in assignment.get("session_ids", [])]
        if not assigned or any(
            session is None or session.status != "graded" or session.submitted_at is None
            or session.submitted_at >= before or session.id not in results_by_session
            for session in assigned
        ):
            continue
        assignment_ids.add(assignment_id)
        session_ids.update(session.id for session in assigned if session is not None)
    for session in sessions_by_id.values():
        if session.course_id not in course_ids or session.id in session_ids or (session.assigned and session.assignment_id):
            continue
        result = results_by_session.get(session.id)
        if result and result.status == "graded" and session.submitted_at and session.submitted_at < before:
            session_ids.add(session.id)
    return session_ids, assignment_ids


def _before(value: str | None, before: datetime) -> bool:
    if not value:
        return False
    try:
        at = datetime.fromisoformat(value)
    except ValueError:
        return False
    if at.tzinfo is None:
        at = at.replace(tzinfo=timezone.utc)
    return at < before


def _cell_value(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, (dict, list, tuple, set)):
        value = json.dumps(value, ensure_ascii=False, default=str)
    if isinstance(value, str):
        if value.lstrip().startswith(("=", "+", "-", "@")):
            value = "'" + value
        return value[:32767]
    return value
